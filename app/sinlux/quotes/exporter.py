# -*- coding: utf-8 -*-
"""报价单输出：PDF (reportlab)、Excel (openpyxl，金额用公式)、WhatsApp 文案。"""
import io
import os
import re
from xml.sax.saxutils import escape

from ..core.util import ApiError
from .settings import render_template

BRAND = '#4F8468'
_FONTS = {}


def _fonts():
    """探测中文字体：微软雅黑 > 黑体 > 宋体 > Linux 常见 CJK 字体；都没有则用 reportlab 自带的 CID 宋体（阅读器侧渲染）。"""
    if _FONTS:
        return _FONTS['regular'], _FONTS['bold']
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular = [('MSYaHei', r'C:\Windows\Fonts\msyh.ttc', 0), ('MSYaHei', r'C:\Windows\Fonts\msyh.ttf', None),
               ('SimHei', r'C:\Windows\Fonts\simhei.ttf', None), ('SimSun', r'C:\Windows\Fonts\simsun.ttc', 0),
               ('Noto', '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 0),
               ('WQY', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', 0), ('WQY', '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc', 0)]
    bold = [('MSYaHeiB', r'C:\Windows\Fonts\msyhbd.ttc', 0), ('MSYaHeiB', r'C:\Windows\Fonts\msyhbd.ttf', None),
            ('SimHeiB', r'C:\Windows\Fonts\simhei.ttf', None)]

    def first(cands):
        for name, path, idx in cands:
            if os.path.exists(path):
                try:
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx) if idx is not None else TTFont(name, path))
                    return name
                except Exception:
                    continue
        return None
    reg = first(regular)
    if not reg:
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        reg = 'STSong-Light'
    _FONTS['regular'], _FONTS['bold'] = reg, first(bold) or reg
    return _FONTS['regular'], _FONTS['bold']


def _para(text, limit=None):
    """用户文本 → reportlab 段落：转义 & < >，换行变 <br/>。"""
    t = (text or '').strip()
    if limit and len(t) > limit:
        t = t[:limit].rstrip() + '…'
    return escape(t).replace('\r\n', '\n').replace('\n', '<br/>')


class QuoteExporter:
    def __init__(self, db, settings, exports_dir, uploads_dir):
        self.db = db
        self.settings = settings
        self.exports_dir = exports_dir
        self.uploads_dir = uploads_dir

    def _image_path(self, product_id):
        """产品主图的缩略图（480px，体积小）；没有缩略图用原图；找不到返回 None。"""
        if not product_id:
            return None
        r = self.db.one('SELECT file, thumb FROM product_images WHERE product_id=? ORDER BY sort_order, id LIMIT 1', (product_id,))
        for name in ((r['thumb'], r['file']) if r else ()):
            if name:
                p = os.path.join(self.uploads_dir, os.path.basename(name))
                if os.path.exists(p):
                    return p
        return None

    @staticmethod
    def _safe_name(quote_no, ext):
        return re.sub(r'[^\w\-.]', '_', quote_no) + '.' + ext

    # ---------- PDF ----------
    def pdf(self, q):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.pdfgen import canvas as rl_canvas
            from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            raise ApiError('缺少 reportlab 库，无法生成 PDF（请确认 app/libs/reportlab 存在）', 500)
        reg, bold = _fonts()
        co = self.settings.company()
        fname = self._safe_name(q['quote_no'], 'pdf')
        path = os.path.join(self.exports_dir, fname)
        os.makedirs(self.exports_dir, exist_ok=True)
        brand = colors.HexColor(BRAND)
        S = lambda size=9, **kw: ParagraphStyle('s', fontName=reg, fontSize=size, leading=size * 1.35, **kw)

        class Numbered(rl_canvas.Canvas):          # 页脚：公司联系方式 + Page x / y
            def __init__(self, *a, **kw):
                super().__init__(*a, **kw)
                self._saved = []

            def showPage(self):
                self._saved.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                total = len(self._saved)
                for st in self._saved:
                    self.__dict__.update(st)
                    self.setFont(reg, 7.5)
                    self.setFillColor(colors.HexColor('#888888'))
                    self.drawString(14 * mm, 9 * mm, ' | '.join(b for b in (co['name'], co['email'], co['phone']) if b))
                    self.drawRightString(A4[0] - 14 * mm, 9 * mm, '%s    Page %d / %d' % (q['quote_no'], self._pageNumber, total))
                    super().showPage()
                super().save()

        doc = SimpleDocTemplate(path, pagesize=A4, topMargin=15 * mm, bottomMargin=18 * mm, leftMargin=14 * mm, rightMargin=14 * mm,
                                title='Quotation ' + q['quote_no'], author=co['name'])
        story = [Paragraph(escape(co['name']), ParagraphStyle('t', fontName=bold, fontSize=21, leading=26, textColor=brand))]
        bits = [b for b in (co['address'], co['email'], co['phone']) if b]
        if bits:
            story.append(Paragraph(escape(' | '.join(bits)), S(8.5, textColor=colors.HexColor('#666666'))))
        story += [Spacer(1, 5 * mm), Paragraph('QUOTATION', ParagraphStyle('q', fontName=bold, fontSize=14, leading=18)), Spacer(1, 2 * mm)]
        left = ['<b>To:</b> ' + _para(q['company'] or q['customer_name'])]
        if q['customer_name'] and q['company']:
            left.append('<b>Attn:</b> ' + _para(q['customer_name']))
        for label, key in (('Address', 'address'), ('Country', 'country')):
            if q.get(key):
                left.append('<b>%s:</b> %s' % (label, _para(q[key])))
        right = ['<b>Quote No.:</b> ' + escape(q['quote_no']), '<b>Date:</b> ' + escape(q['created_at'][:10]),
                 '<b>Valid until:</b> %s (%d days)' % (escape(q['valid_until']), q['valid_days']), '<b>Currency:</b> ' + escape(q['currency'])]
        meta = Table([[Paragraph('<br/>'.join(left), S()), Paragraph('<br/>'.join(right), S())]], colWidths=[100 * mm, 82 * mm])
        meta.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LINEBELOW', (0, 0), (-1, 0), 0.8, brand), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
        story += [meta, Spacer(1, 4 * mm)]

        cur = q['currency']
        rows = [['#', 'Image', 'SKU', 'Description', 'Qty', 'Unit', 'Price (%s)' % cur, 'Amount (%s)' % cur]]
        for n, i in enumerate(q['items'], 1):
            desc = '<b>%s</b>' % _para(i['name'] or i['sku'])
            if i['spec']:
                desc += '<br/><font size="7.5" color="#555555">%s</font>' % _para(i['spec'], 1500)
            if i['remark']:
                desc += '<br/><font size="7" color="#888888">%s</font>' % _para(i['remark'], 300)
            img = ''
            ip = self._image_path(i['product_id'])
            if ip:
                try:
                    img = Image(ip, width=17 * mm, height=17 * mm, kind='proportional')
                except Exception:
                    img = ''
            rows.append([str(n), img, Paragraph(escape(i['sku'] or ''), S(8)), Paragraph(desc, S(8.5)), '%g' % i['quantity'], i['unit'],
                         '{:,.2f}'.format(i['unit_price']), '{:,.2f}'.format(i['amount'])])
        rows.append(['', '', '', '', '', '', 'TOTAL', '{:,.2f}'.format(q['total'])])
        t = Table(rows, colWidths=[7 * mm, 20 * mm, 24 * mm, 63 * mm, 12 * mm, 11 * mm, 22 * mm, 23 * mm], repeatRows=1, splitInRow=1)
        t.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), reg), ('FONTSIZE', (0, 0), (-1, -1), 8), ('FONTNAME', (0, 0), (-1, 0), bold),
            ('BACKGROUND', (0, 0), (-1, 0), brand), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -2), 0.4, colors.HexColor('#D8D8D8')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (4, 0), (-1, -1), 'RIGHT'), ('ALIGN', (0, 0), (0, -1), 'CENTER'), ('ALIGN', (1, 0), (1, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F6F9F6')]),
            ('FONTNAME', (6, -1), (-1, -1), bold), ('FONTSIZE', (6, -1), (-1, -1), 10), ('TEXTCOLOR', (6, -1), (-1, -1), brand),
            ('LINEABOVE', (0, -1), (-1, -1), 1, brand), ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]))
        story += [t, Spacer(1, 5 * mm)]
        for label, key in (('Lead time', 'lead_time'), ('Payment terms', 'payment_terms'), ('Trade terms', 'shipping_terms'), ('Notes', 'notes')):
            if q.get(key):
                story.append(Paragraph('<b>%s:</b> %s' % (label, _para(q[key])), S(9)))
        story += [Spacer(1, 7 * mm), Paragraph('Thank you for your business!', S(9.5, textColor=brand))]
        doc.build(story, canvasmaker=Numbered)
        return fname

    # ---------- Excel ----------
    def excel(self, q):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        co = self.settings.company()
        fname = self._safe_name(q['quote_no'], 'xlsx')
        os.makedirs(self.exports_dir, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Quotation'
        thin = Side(style='thin', color='D8D8D8')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        green = '4F8468'

        def put(ref, value, **font):
            c = ws[ref]
            c.value = value
            if isinstance(value, str) and value.startswith('='):
                c.data_type = 's'                  # 用户文本里以 = 开头的，当文本存，不当公式
            if font:
                c.font = Font(name='Arial', **font)
            return c
        ws.merge_cells('A1:I1'); put('A1', co['name'], bold=True, size=18, color=green)
        ws.merge_cells('A2:I2'); put('A2', ' | '.join(b for b in (co['address'], co['email'], co['phone']) if b), size=9, color='666666')
        ws.merge_cells('A3:I3'); put('A3', 'QUOTATION', bold=True, size=13)
        put('A5', 'To: %s' % (q['company'] or q['customer_name'] or ''), bold=True)
        put('A6', 'Attn: %s' % (q['customer_name'] or '')); put('A7', 'Country: %s' % (q['country'] or ''))
        put('G5', 'Quote No.: %s' % q['quote_no']); put('G6', 'Date: %s' % q['created_at'][:10])
        put('G7', 'Valid until: %s (%d days)' % (q['valid_until'], q['valid_days']))
        hdr = 9
        heads = ['#', 'Image', 'SKU', 'Description', 'Specification', 'Qty', 'Unit', 'Price (%s)' % q['currency'], 'Amount (%s)' % q['currency']]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(row=hdr, column=c, value=h)
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', start_color=green)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        r = hdr + 1
        try:
            from openpyxl.drawing.image import Image as XImage
            from PIL import Image as PImage
        except Exception:
            XImage = PImage = None
        for n, i in enumerate(q['items'], 1):
            vals = [n, None, i['sku'], i['name'], i['spec'], i['quantity'], i['unit'], i['unit_price'], None]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                if isinstance(v, str) and v.startswith('='):
                    cell.data_type = 's'
                cell.border = border
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=c in (4, 5))
            ws.cell(row=r, column=9, value='=ROUND(F%d*H%d,2)' % (r, r)).border = border      # 逐行四舍五入，合计与 PDF 一致
            for c in (8, 9):
                ws.cell(row=r, column=c).number_format = '#,##0.00'
            ws.row_dimensions[r].height = 52
            ip = self._image_path(i['product_id'])
            if ip and XImage:
                try:
                    buf = io.BytesIO()
                    with PImage.open(ip) as im:
                        im = im.convert('RGB'); im.thumbnail((120, 120)); im.save(buf, 'PNG')
                    buf.seek(0)
                    xi = XImage(buf); xi.width, xi.height = 60, 60
                    ws.add_image(xi, 'B%d' % r)
                except Exception:
                    pass
            r += 1
        ws.cell(row=r, column=8, value='TOTAL').font = Font(name='Arial', bold=True)
        tc = ws.cell(row=r, column=9, value='=SUM(I%d:I%d)' % (hdr + 1, r - 1))
        tc.font = Font(name='Arial', bold=True, color=green); tc.number_format = '#,##0.00'; tc.border = border
        r += 2
        for label, key in (('Lead time', 'lead_time'), ('Payment terms', 'payment_terms'), ('Trade terms', 'shipping_terms'), ('Notes', 'notes')):
            if q.get(key):
                c = ws.cell(row=r, column=1, value='%s: %s' % (label, q[key]))
                if q[key].startswith('='):
                    c.data_type = 's'
                r += 1
        for col, w in zip('ABCDEFGHI', [5, 11, 16, 34, 40, 8, 8, 14, 14]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A%d' % (hdr + 1)
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.save(os.path.join(self.exports_dir, fname))
        return fname

    # ---------- WhatsApp ----------
    def whatsapp(self, q):
        lines = []
        for i in q['items']:
            spec = re.sub(r'\s+', ' ', i['spec'] or '').strip()
            spec = ' (%s)' % (spec[:80] + ('…' if len(spec) > 80 else '')) if spec else ''
            lines.append('- %s %s%s: %g %s × %s %g = %s %.2f' % (i['sku'], i['name'], spec, i['quantity'], i['unit'], q['currency'],
                                                               i['unit_price'], q['currency'], i['amount']))
        values = {'customer_name': q['customer_name'] or q['company'] or 'there', 'company': q['company'] or '', 'quote_no': q['quote_no'],
                  'items': '\n'.join(lines), 'currency': q['currency'], 'total': '%.2f' % q['total'], 'lead_time': q['lead_time'] or 'TBD',
                  'payment_terms': q['payment_terms'] or 'TBD', 'valid_days': q['valid_days'], 'valid_until': q['valid_until'],
                  'shipping_terms': q['shipping_terms'] or 'TBD', 'notes': q['notes'] or '', 'sender': self.settings.company()['name']}
        return render_template(self.settings.template(), values)
