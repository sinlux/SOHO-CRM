# -*- coding: utf-8 -*-
"""把 .xls / .xlsx 读成同一种「格子」：grid.cell(行, 列)，行列都从 0 开始，值是 str / int / float / datetime / None。
PI、供应商报价单这类文件通常只有几十到几百行，整张读进内存没问题；大产品清单仍走产品导入向导的流式读取。"""
import datetime as _dt
import os

MAX_ROWS = 6000
MAX_COLS = 80


class Grid:
    def __init__(self, rows, sheet, sheets, path):
        self.rows, self.sheet, self.sheets, self.path = rows, sheet, sheets, path
        self.nrows = len(rows)
        self.ncols = max((len(r) for r in rows), default=0)

    def cell(self, r, c):
        if 0 <= r < self.nrows and 0 <= c < len(self.rows[r]):
            return self.rows[r][c]
        return None

    def text(self, r, c):
        v = self.cell(r, c)
        if v is None:
            return ''
        if isinstance(v, float) and v == int(v):
            return str(int(v))
        if isinstance(v, (_dt.datetime, _dt.date)):
            return v.strftime('%Y-%m-%d')
        return str(v).strip()

    def row_texts(self, r):
        return [self.text(r, c) for c in range(self.ncols)]


def _norm(v):
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def read_grid(path, filename, sheet=None):
    """按扩展名读取。返回 Grid；sheet 为空时选第一张有内容的表。"""
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.xls':
        return _read_xls(path, sheet)
    if ext in ('.xlsx', '.xlsm'):
        return _read_xlsx(path, sheet)
    raise ValueError('只支持 .xls / .xlsx 文件')


def _read_xls(path, sheet):
    import xlrd
    wb = xlrd.open_workbook(path)
    names = wb.sheet_names()
    pick = sheet if sheet in names else next((n for n in names if wb.sheet_by_name(n).nrows), names[0])
    ws = wb.sheet_by_name(pick)
    rows = []
    for r in range(min(ws.nrows, MAX_ROWS)):
        row = []
        for c in range(min(ws.ncols, MAX_COLS)):
            t, v = ws.cell_type(r, c), ws.cell_value(r, c)
            if t == xlrd.XL_CELL_DATE:
                try:
                    v = xlrd.xldate_as_datetime(v, wb.datemode)
                except Exception:
                    pass
            elif t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                v = None
            elif t == xlrd.XL_CELL_BOOLEAN:
                v = bool(v)
            elif t == xlrd.XL_CELL_ERROR:
                v = None
            row.append(_norm(v))
        rows.append(row)
    return Grid(rows, pick, names, path)


def _read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        order = [sheet] if sheet in names else names
        for n in order:
            rows = []
            for i, row in enumerate(wb[n].iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append([_norm(v) for v in row[:MAX_COLS]])
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            if rows or sheet in names:
                return Grid(rows, n, names, path)
        return Grid([], names[0], names, path)
    finally:
        wb.close()
