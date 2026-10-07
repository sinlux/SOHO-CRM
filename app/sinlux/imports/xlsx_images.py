# -*- coding: utf-8 -*-
"""大 Excel 的读取辅助：表头/行预览用 openpyxl 只读模式；图片不加载整个工作簿，
而是把 xlsx 当 zip 包，解析 drawing XML 里的锚点（oneCellAnchor / twoCellAnchor）找到图片所在行，
Target 路径同时支持绝对（/xl/...）和相对（../media/...）写法；每行只取第一张图，压缩成长边 800px 的 JPEG(q82)。"""
import io
import os
import zipfile
from xml.etree import ElementTree as ET

NS = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def _cell(c):
    if c is None:
        return ''
    if isinstance(c, float) and c == int(c) and abs(c) < 1e15:
        return int(c)
    return c


def list_sheets(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def inspect_sheet(path, sheet, header_row=1, preview=8):
    """返回 {headers, preview_raw, total_rows}。headers 空列用「列N」代替；重复表头自动加 (2)(3) 区分。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        headers, raw, total = [], [], 0
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            total = i
            if i == header_row:
                seen = {}
                for j, c in enumerate(row):
                    h = str(c).strip() if c is not None and str(c).strip() else '列%d' % (j + 1)
                    seen[h] = seen.get(h, 0) + 1
                    headers.append(h if seen[h] == 1 else '%s(%d)' % (h, seen[h]))
            if i <= max(8, header_row + preview) and len(raw) < 20:
                raw.append([_cell(c) for c in row])
        return {'headers': headers, 'preview_raw': raw, 'total_rows': total}
    finally:
        wb.close()


def iter_rows(path, sheet, header_row):
    """逐行产出 (行号, 值列表)，行号与 Excel 一致（1 开始），跳过表头及以上和全空行。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for i, row in enumerate(wb[sheet].iter_rows(values_only=True), 1):
            if i <= header_row:
                continue
            if all(c is None or str(c).strip() == '' for c in row):
                continue
            yield i, [_cell(c) for c in row]
    finally:
        wb.close()


def _resolve(source_part, target):
    if target.startswith('/'):
        return target.lstrip('/')
    return os.path.normpath(os.path.join(os.path.dirname(source_part), target)).replace('\\', '/')


def _sheet_part(zf, sheet):
    wb = ET.fromstring(zf.read('xl/workbook.xml'))
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))}
    for sh in wb.find('main:sheets', NS):
        if sh.get('name') == sheet:
            t = rels.get(sh.get('{%s}id' % NS['r']))
            if t:
                path = _resolve('xl/workbook.xml', t)
                parts = path.split('/')
                return path, '/'.join(parts[:-1] + ['_rels', parts[-1] + '.rels'])
    return None, None


def extract_row_images(path, sheet, dest_dir, max_side=800, quality=82):
    """返回 {行号: 文件名}（文件存在 dest_dir）。出错的单张图跳过，不影响其它。"""
    from PIL import Image
    out = {}
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        sp, rp = _sheet_part(zf, sheet)
        if not sp or rp not in names:
            return out
        drawing = None
        for rel in ET.fromstring(zf.read(rp)):
            if rel.get('Type', '').endswith('/drawing'):
                drawing = _resolve(sp, rel.get('Target'))
                break
        if not drawing or drawing not in names:
            return out
        dxml = ET.fromstring(zf.read(drawing))
        drp = os.path.dirname(drawing) + '/_rels/' + os.path.basename(drawing) + '.rels'
        media = {}
        if drp in names:
            media = {r.get('Id'): _resolve(drawing, r.get('Target')) for r in ET.fromstring(zf.read(drp))}
        os.makedirs(dest_dir, exist_ok=True)
        for tag in ('twoCellAnchor', 'oneCellAnchor'):
            for anchor in dxml.findall('xdr:' + tag, NS):
                frm = anchor.find('xdr:from', NS)
                blip = anchor.find('.//a:blip', NS)
                if frm is None or blip is None or frm.find('xdr:row', NS) is None:
                    continue
                row = int(frm.find('xdr:row', NS).text) + 1
                part = media.get(blip.get('{%s}embed' % NS['r']))
                if row in out or not part or part not in names:
                    continue
                try:
                    with Image.open(io.BytesIO(zf.read(part))) as im:
                        im.load()
                        if im.mode in ('RGBA', 'LA', 'P'):
                            im = im.convert('RGBA')
                            bg = Image.new('RGB', im.size, (255, 255, 255))
                            bg.paste(im, mask=im.split()[3])
                            im = bg
                        else:
                            im = im.convert('RGB')
                        im.thumbnail((max_side, max_side))
                        fname = 'row%d.jpg' % row
                        im.save(os.path.join(dest_dir, fname), 'JPEG', quality=quality)
                        out[row] = fname
                except Exception:
                    continue
    return out
