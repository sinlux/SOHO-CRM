# -*- coding: utf-8 -*-
"""客户 Excel 导入（先预览、确认后写入）与导出。

导入会话 = data/import_tmp/customers/<sid>/ 目录里的原始文件，落在磁盘上，
程序重启后仍可继续（旧版放内存，重启即丢）。预览和写入用同一套解析逻辑，所以所见即所得。
"""
import os
import re
import shutil
import time
import uuid

from ..core.util import ApiError, clean_multi
from . import service as cs

COL_MAP = {'lv': 'lv', 'country': 'country', 'name': 'name', 'company': 'company', 'website': 'website',
           'email': 'emails', 'emails': 'emails', 'main business': 'main_business', 'address': 'address',
           'whatsapp': 'whatsapp', 'linkedin': 'linkedin', 'facebook': 'facebook', 'remark': 'remark'}
MAX_UPLOAD = 100 * 1024 * 1024
SID_RE = re.compile(r'^[0-9a-f]{12}$')


class CustomerImporter:
    def __init__(self, db, customers, tmp_root):
        self.db = db
        self.customers = customers
        self.tmp_root = tmp_root

    # ---------- 会话 ----------
    def new_session_path(self, filename):
        sid = uuid.uuid4().hex[:12]
        d = os.path.join(self.tmp_root, sid)
        os.makedirs(d, exist_ok=True)
        return sid, os.path.join(d, 'upload.xlsx'), filename

    def _session_file(self, sid):
        if not SID_RE.match(sid or ''):
            raise ApiError('导入会话无效')
        p = os.path.join(self.tmp_root, sid, 'upload.xlsx')
        if not os.path.isfile(p):
            raise ApiError('导入会话不存在或已清理，请重新上传文件', 404)
        return p

    def cleanup_old(self, max_age_hours=24):
        if not os.path.isdir(self.tmp_root):
            return
        cutoff = time.time() - max_age_hours * 3600
        for n in os.listdir(self.tmp_root):
            p = os.path.join(self.tmp_root, n)
            if os.path.getmtime(p) < cutoff:
                shutil.rmtree(p, ignore_errors=True)

    # ---------- 解析 ----------
    def _parse(self, path, remap_lv):
        """纯解析 + 查重，不写库。返回 (rows, problems, stats)。"""
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise ApiError('无法读取 Excel 文件（需要 .xlsx 格式）: %s' % e)
        try:
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                raise ApiError('表格为空')
            idx = {}
            for i, h in enumerate(header):
                if h is not None and str(h).strip().lower() in COL_MAP:
                    idx.setdefault(COL_MAP[str(h).strip().lower()], i)
            if 'company' not in idx and 'emails' not in idx:
                raise ApiError('表头无法识别，读到的表头是: %s' % [h for h in header if h is not None])
            raw_rows = list(enumerate(it, start=2))
            sheet = wb.sheetnames[0]
        finally:
            wb.close()

        known_companies, known_emails = set(), set()
        for c in self.db.query('SELECT company, emails FROM customers'):
            if (c['company'] or '').strip():
                known_companies.add(c['company'].strip().lower())
            known_emails.update(e.lower() for e in (c['emails'] or '').split())

        rows, problems, dup_n = [], [], 0
        for rownum, row in raw_rows:
            if row is None or all(v is None or str(v).strip() == '' for v in row):
                continue

            def cell(f, row=row):
                i = idx.get(f)
                v = row[i] if i is not None and i < len(row) else None
                return str(v).strip() if v is not None else ''

            d = {'country': cell('country'), 'name': cell('name'), 'company': cell('company'),
                 'website': clean_multi(cell('website')), 'emails': clean_multi(cell('emails')),
                 'main_business': cell('main_business'), 'address': cell('address'),
                 'whatsapp': cell('whatsapp'), 'linkedin': clean_multi(cell('linkedin')),
                 'facebook': clean_multi(cell('facebook'))}
            lv, lv_raw = None, cell('lv')
            if lv_raw:
                try:
                    n = int(float(lv_raw))
                    if not 1 <= n <= 6 or float(lv_raw) != n:
                        raise ValueError
                    lv = 7 - n if remap_lv else n
                except ValueError:
                    problems.append('第%d行 LV 值无效: "%s"（需要 1-6，已置空）' % (rownum, lv_raw))
            d['lv'] = lv
            remark = cell('remark')

            who = d['company'] or d['emails'] or d['name']
            if not (d['company'] or d['name'] or d['emails']):
                problems.append('第%d行 公司/联系人/邮箱全为空，跳过' % rownum)
                rows.append({'row': rownum, 'action': 'skip', 'reason': '空行', 'data': d})
                continue
            is_dup = bool(d['company']) and d['company'].lower() in known_companies
            is_dup = is_dup or any(e.lower() in known_emails for e in d['emails'].split())
            if is_dup:
                dup_n += 1
                problems.append('第%d行 与已有客户重复（%s），跳过' % (rownum, who))
                rows.append({'row': rownum, 'action': 'skip', 'reason': '重复', 'data': d})
                continue
            if d['company']:
                known_companies.add(d['company'].lower())
            known_emails.update(e.lower() for e in d['emails'].split())
            rows.append({'row': rownum, 'action': 'import', 'data': d, 'remark': remark})
        stats = {'sheet': sheet, 'total_rows': len(rows),
                 'to_import': sum(1 for r in rows if r['action'] == 'import'),
                 'duplicates': dup_n,
                 'skipped': sum(1 for r in rows if r['action'] == 'skip')}
        return rows, problems, stats

    def preview(self, sid, remap_lv=False, sample=30):
        rows, problems, stats = self._parse(self._session_file(sid), remap_lv)
        stats.update(session_id=sid, remap_lv=bool(remap_lv), problems=problems,
                     sample=[r for r in rows if r['action'] == 'import'][:sample])
        return stats

    def apply(self, sid, remap_lv=False):
        rows, problems, stats = self._parse(self._session_file(sid), remap_lv)
        imported = 0
        with self.db.tx():              # 整批一个事务：中途出错全部回滚，不会导入一半
            for r in rows:
                if r['action'] != 'import':
                    continue
                cid = self.customers.create(r['data'])
                if r['remark']:
                    self.customers.add_system_note(cid, '【导入时的Remark】' + r['remark'])
                imported += 1
        shutil.rmtree(os.path.join(self.tmp_root, sid), ignore_errors=True)
        return {'imported': imported, 'skipped': stats['skipped'], 'problems': problems}


# ---------- 导出 ----------
EXPORT_COLS = [('lv', 'LV'), ('stage', '阶段'), ('country', '国家'), ('name', '联系人'), ('company', '公司'),
               ('website', '网站'), ('emails', '邮箱'), ('whatsapp', 'WhatsApp'), ('linkedin', 'LinkedIn'),
               ('facebook', 'Facebook'), ('other_social', '其他社媒'), ('main_business', '主营业务'),
               ('address', '地址'), ('company_size', '公司规模'), ('founded_year', '成立年份'),
               ('customer_type', '客户类型'), ('certifications', '认证'), ('ai_summary', 'AI摘要'),
               ('created_at', '创建时间')]


def export_customers(db, export_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    os.makedirs(export_dir, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = '客户'
    for c, (_, label) in enumerate(EXPORT_COLS, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', start_color='1F2937')
    for r, row in enumerate(db.query('SELECT * FROM customers ORDER BY COALESCE(lv,0) DESC, id'), 2):
        for c, (key, _) in enumerate(EXPORT_COLS, 1):
            v = row.get(key)
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, str) and v.startswith('='):
                cell.data_type = 's'   # 当文本存，防止客户资料里的 "=..." 被 Excel 当公式执行
    widths = [5, 8, 10, 14, 26, 22, 30, 16, 24, 24, 20, 30, 30, 10, 8, 10, 14, 40, 18]
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
    name = time.strftime('客户导出_%Y%m%d_%H%M%S') + '.xlsx'
    wb.save(os.path.join(export_dir, name))
    return name
