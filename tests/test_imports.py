# -*- coding: utf-8 -*-
"""第四批：Excel 产品批量导入（5 步向导）+ PI 导入。"""
import io
import os
import shutil
import tempfile
import unittest
import urllib.parse
import zipfile

import xlwt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

from helpers import AppTestCase
from imgutil import PNG_RED_BOX, png_bytes


def make_products_xlsx(rows, header_row=2, images=None, two_cell=False, title='SINLUX 货盘 2026'):
    """rows: [[序号, 图片, 款号, 品名, 货盘价, 币种, MOQ, 供应商, 瓦数, 颜色]]。images: {excel行号: png bytes}。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '货盘'
    ws.cell(row=1, column=1, value=title)
    for c, h in enumerate(['序号', '图片', '款号', '品名', '货盘价', '币种', 'MOQ', '供应商', '瓦数', '颜色'], 1):
        ws.cell(row=header_row, column=c, value=h)
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            ws.cell(row=header_row + 1 + i, column=c, value=v)
    for rn, data in (images or {}).items():
        img = XImage(io.BytesIO(data))
        if two_cell:
            img.anchor = TwoCellAnchor(editAs='oneCell', _from=AnchorMarker(col=1, row=rn - 1), to=AnchorMarker(col=2, row=rn))
        else:
            img.anchor = 'B%d' % rn
        ws.add_image(img)
    ws2 = wb.create_sheet('空表')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportBase(AppTestCase):
    def upload(self, data, name='货盘.xlsx'):
        return self.c.j('POST', '/api/import/products/upload', raw=data, headers={'X-Filename': urllib.parse.quote(name)})

    def wizard(self, data, cat='lighting', mapping_edit=None, header_row=2, **pv):
        st, r = self.upload(data)
        self.assertEqual(st, 200, r)
        sid = r['session_id']
        st, s = self.c.post('/api/import/products/%s/sheet' % sid, {'sheet': '货盘', 'header_row': header_row})
        self.assertEqual(st, 200, s)
        st, c = self.c.post('/api/import/products/%s/category' % sid, {'category_id': self.cat_id(cat)})
        self.assertEqual(st, 200, c)
        mapping = dict(c['suggested_mapping'])
        mapping.update(mapping_edit or {})
        body = {'mapping': mapping, 'currency': 'CNY', 'profit_rate': 0.3}
        body.update(pv)
        st, p = self.c.post('/api/import/products/%s/preview' % sid, body)
        return sid, c, p, st


class TestProductImport(ImportBase):
    ROWS = [[1, None, 'IMP-001', 'Pendant A', 25.5, 'CNY', 500, '甲厂', '12W', 'Black'],
            [2, None, 'IMP-002', 'Pendant B', '¥30', 'RMB', 300, '乙厂', '15W', 'White'],
            [3, None, 'IMP-003', 'Wall lamp', 4.2, 'USD', None, None, '6W', None]]

    def test_full_wizard_with_images_and_mapping_guess(self):
        data = make_products_xlsx(self.ROWS, images={3: PNG_RED_BOX, 4: PNG_RED_BOX})
        sid, c, p, st = self.wizard(data)
        self.assertEqual(st, 200, p)
        sm = c['suggested_mapping']
        self.assertEqual((sm['序号'], sm['图片'], sm['款号'], sm['品名'], sm['货盘价'], sm['币种'], sm['MOQ'], sm['供应商']),
                         ('__ignore', '__ignore', '__sku', '__name', '__cost', '__currency', '__moq', '__supplier'))
        self.assertEqual(sm['颜色'], '__spec')
        self.assertEqual((p['count'], p['with_image'], p['blocked'], p['will_update'], p['warned']), (3, 2, 0, 0, 0))
        r0, r1, r2 = p['rows']
        self.assertEqual((r0['row'], r0['sku'], r0['cost'], r0['currency'], r0['moq'], r0['supplier']), (3, 'IMP-001', 25.5, 'CNY', 500, '甲厂'))
        self.assertEqual((r1['cost'], r1['currency']), (30.0, 'CNY'))                       # ¥30 / RMB 都认得
        self.assertEqual((r2['currency'], r2['moq']), ('USD', None))
        self.assertIn('颜色: Black', r0['spec_text'])
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM products WHERE sku LIKE 'IMP-%'"), 0)      # 预览不写库
        st, raw, h = self.c.call('GET', '/api/import/products/%s/image/%s' % (sid, r0['image']))
        self.assertEqual((st, h['Content-Type'].startswith('image/')), (200, True))
        self.assertEqual(self.c.call('GET', '/api/import/products/%s/image/..%%2f..%%2fupload.xlsx' % sid)[0], 404)
        st, res = self.c.post('/api/import/products/%s/apply' % sid, {'skip_rows': [5]})
        self.assertEqual((st, res['created'], res['updated'], res['skipped'], res['failed']), (200, 2, 0, 1, []))
        db = self.ctx.db
        self.assertIsNone(db.one("SELECT 1 FROM products WHERE sku='IMP-003'"))
        p1 = db.one("SELECT * FROM products WHERE sku='IMP-001'")
        self.assertEqual((p1['cost'], p1['cost_currency'], p1['moq'], p1['supplier'], p1['profit_rate'], p1['name']), (25.5, 'CNY', 500, '甲厂', 0.3, 'Pendant A'))
        self.assertEqual(p1['suggested_price'], round(25.5 * 0.138 * 1.3, 2))
        self.assertIn('颜色: Black', p1['spec_text'])
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM product_images WHERE product_id=?', (p1['id'],)), 1)       # 图片入库且已规范化
        self.assertEqual(db.scalar("SELECT normalized FROM product_images WHERE product_id=?", (p1['id'],)), 1)
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM price_history WHERE product_id=? AND price_type='cost'", (p1['id'],)), 1)
        self.assertEqual(self.c.get('/api/import/products/%s/rows' % sid)[0], 404)                  # 会话用完即清
        self.assertFalse(os.path.exists(os.path.join(self.ctx.import_tmp, 'products', sid)))

    def test_structured_field_mapping_and_update_existing(self):
        fields = self.c.get('/api/categories/%d/fields' % self.cat_id('lighting'))[1]['fields']
        watt = next(f for f in fields if f['label'] in ('功率', '瓦数', 'Wattage', 'Power') or 'watt' in f['key'].lower() or 'power' in f['key'].lower())
        pid = self.new_product(sku='UPD-1', name='Old name', cost=10, cost_currency='CNY', remark='keep me', supplier='旧厂')
        up = self.c.j('POST', '/api/images/upload', raw=PNG_RED_BOX)[1]['image']
        self.c.put('/api/products/%d' % pid, {'images': [up['id']]})
        rows = [[1, None, 'upd-1', 'New name', 12, 'CNY', None, None, '9W', 'Gold']]
        sid, c, p, st = self.wizard(make_products_xlsx(rows, images={3: png_bytes(200, 200, box=(50, 50, 150, 150), box_color=(0, 0, 200))}),
                                    mapping_edit={'瓦数': 'f:%d' % watt['id']})
        self.assertEqual((st, p['will_update']), (200, 1))                                          # SKU 大小写不敏感
        res = self.c.post('/api/import/products/%s/apply' % sid, {})[1]
        self.assertEqual((res['created'], res['updated'], res['failed']), (0, 1, []))
        d = self.c.get('/api/products/%d' % pid)[1]['product']
        self.assertEqual((d['sku'], d['name'], d['cost'], d['remark'], d['supplier']), ('UPD-1', 'New name', 12.0, 'keep me', '旧厂'))   # 没映射的不动
        self.assertEqual(d['field_values'][str(watt['id'])], '9W')
        self.assertEqual(len(d['images']), 1)                                                       # 已有图片不被覆盖
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM price_history WHERE product_id=? AND price_type='cost'", (pid,)), 2)
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM products WHERE sku='UPD-1' COLLATE NOCASE"), 1)

    def test_warnings_blocked_rows_and_dup_in_file(self):
        rows = [[1, None, 'W-1', 'ok', 'abc', 'EUR', 'x', None, None, None],
                [2, None, None, 'no sku', 5, 'CNY', None, None, None, None],
                [3, None, 'W-1', 'dup', 7, 'CNY', None, None, None, None]]
        sid, c, p, st = self.wizard(make_products_xlsx(rows))
        self.assertEqual((st, p['blocked'], p['warned']), (200, 1, 3))
        w = p['rows'][0]['warnings']
        self.assertTrue(any('成本' in x for x in w) and any('币种' in x for x in w) and any('MOQ' in x for x in w), w)
        self.assertEqual(p['rows'][0]['currency'], 'CNY')                                            # EUR 不支持，退回统一币种
        self.assertTrue(any('重复' in x for x in p['rows'][2]['warnings']))
        res = self.c.post('/api/import/products/%s/apply' % sid, {})[1]
        self.assertEqual((res['created'], res['updated'], [f['row'] for f in res['failed']]), (1, 1, [4]))
        pr = self.ctx.db.one("SELECT * FROM products WHERE sku='W-1'")
        self.assertEqual((pr['name'], pr['cost']), ('dup', 7.0))                                    # 后者覆盖前者

    def test_validation_errors(self):
        self.assertEqual(self.upload(b'x', 'a.xls')[0], 400)
        st, r = self.upload(b'not a zip', 'bad.xlsx')
        self.assertEqual(st, 400)
        self.assertEqual(self.c.post('/api/import/products/nosuchsession/sheet', {'sheet': 'x'})[0], 404)
        st, r = self.upload(make_products_xlsx(self.ROWS))
        sid = r['session_id']
        self.assertEqual(r['sheets'], ['货盘', '空表'])
        for body in ({'sheet': '没有这张表'}, {'sheet': '货盘', 'header_row': 0}, {'sheet': '货盘', 'header_row': 'x'}, {'sheet': '空表', 'header_row': 1}):
            self.assertEqual(self.c.post('/api/import/products/%s/sheet' % sid, body)[0], 400, body)
        self.assertEqual(self.c.post('/api/import/products/%s/preview' % sid, {'mapping': {}})[0], 400)       # 还没选表/类目
        self.c.post('/api/import/products/%s/sheet' % sid, {'sheet': '货盘', 'header_row': 2})
        self.assertEqual(self.c.post('/api/import/products/%s/category' % sid, {'category_id': 99999})[0], 404)
        c = self.c.post('/api/import/products/%s/category' % sid, {'category_id': self.cat_id('lighting')})[1]
        m = dict(c['suggested_mapping'])
        for bad in ({'mapping': {k: ('__ignore' if v == '__sku' else v) for k, v in m.items()}},          # 没有 SKU 列
                    {'mapping': {**m, '品名': '__nope'}}, {'mapping': {**m, '不存在的列': '__sku'}},
                    {'mapping': m, 'currency': 'EUR'}, {'mapping': m, 'profit_rate': 'x'}, {'mapping': m, 'profit_rate': 99}):
            self.assertEqual(self.c.post('/api/import/products/%s/preview' % sid, bad)[0], 400, bad)
        self.assertEqual(self.c.post('/api/import/products/%s/apply' % sid, {})[0], 400)                   # 没预览不能写入
        self.assertEqual(self.c.delete('/api/import/products/%s' % sid)[0], 200)
        self.assertEqual(self.c.post('/api/import/products/%s/apply' % sid, {})[0], 404)

    def test_preview_paging_and_no_images(self):
        rows = [[i, None, 'PG-%03d' % i, 'n%d' % i, i, 'CNY', None, None, None, None] for i in range(1, 251)]
        sid, c, p, st = self.wizard(make_products_xlsx(rows), extract_images=False)
        self.assertEqual((p['count'], len(p['rows']), p['with_image']), (250, 200, 0))
        pg = self.c.get('/api/import/products/%s/rows?offset=200&limit=100' % sid)[1]
        self.assertEqual((len(pg['rows']), pg['rows'][0]['sku']), (50, 'PG-201'))
        self.assertEqual(self.c.delete('/api/import/products/%s' % sid)[0], 200)

    def test_header_row_other_than_one_and_blank_header_cells(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '货盘'
        for r, vals in enumerate([['title'], [], ['款号', None, '款号', '成本'], ['H-1', 'x', 'dupcol', 9]], 1):
            for c, v in enumerate(vals, 1):
                ws.cell(row=r, column=c, value=v)
        buf = io.BytesIO(); wb.save(buf)
        st, r = self.upload(buf.getvalue())
        s = self.c.post('/api/import/products/%s/sheet' % r['session_id'], {'sheet': '货盘', 'header_row': 3})[1]
        self.assertEqual(s['headers'], ['款号', '列2', '款号(2)', '成本'])                            # 空表头、重复表头都有唯一名字
        self.c.delete('/api/import/products/%s' % r['session_id'])


class TestRowImages(unittest.TestCase):
    """图片提取：两种锚点（oneCellAnchor / twoCellAnchor）、绝对和相对 Target 路径。"""

    def extract(self, data, absolute=False):
        from sinlux.imports import xlsx_images as xi
        d = tempfile.mkdtemp()
        try:
            p = os.path.join(d, 'a.xlsx')
            if absolute:
                src = zipfile.ZipFile(io.BytesIO(data))
                out = io.BytesIO()
                with zipfile.ZipFile(out, 'w') as z:
                    for n in src.namelist():
                        b = src.read(n)
                        if n.startswith('xl/drawings/_rels/'):
                            b = b.replace(b'../media/', b'/xl/media/')
                        z.writestr(n, b)
                data = out.getvalue()
            open(p, 'wb').write(data)
            return xi.extract_row_images(p, '货盘', os.path.join(d, 'img')), d
        finally:
            pass

    def test_anchor_types_and_target_styles(self):
        for two in (False, True):
            for absolute in (False, True):
                data = make_products_xlsx([[1, None, 'A', 'a', 1, 'CNY', None, None, None, None]] * 3, images={3: PNG_RED_BOX, 5: PNG_RED_BOX}, two_cell=two)
                got, d = self.extract(data, absolute)
                self.assertEqual(sorted(got), [3, 5], (two, absolute))
                from PIL import Image
                with Image.open(os.path.join(d, 'img', got[3])) as im:
                    self.assertEqual((im.format, max(im.size) <= 800), ('JPEG', True))
                shutil.rmtree(d, ignore_errors=True)

    def test_no_images_and_wrong_sheet(self):
        got, d = self.extract(make_products_xlsx([[1, None, 'A', 'a', 1, 'CNY', None, None, None, None]]))
        self.assertEqual(got, {})
        shutil.rmtree(d, ignore_errors=True)


# ---------------------------------------------------------------- PI
def make_pi(invoice='SL-20260101US', date_serial=46023, buyer=('ACME Hotels Ltd', 'Bob', '1 Main St, Kingston', 'bob@acme-hotels.com', '+1 555 0100'),
            items=None, currency='USD'):
    wb = xlwt.Workbook()
    ws = wb.add_sheet('PI')
    ws.write(1, 0, 'Invoice Number:'); ws.write(1, 2, invoice)
    ws.write(2, 0, 'Issued Date:'); ws.write(2, 2, date_serial)
    for i, (lab, v) in enumerate(zip(['Buyer:', 'Contact Person:', 'Address:', 'Email:', 'Phone:'], buyer)):
        ws.write(5 + i, 0, lab); ws.write(5 + i, 1, v)
    heads = ['Item No.', 'Product Name', 'Parameters', 'Picture', 'Price/PCS (%s)' % currency, 'Quantity', 'Amount', 'rmb', 'kg', 'cbm']
    for c, h in enumerate(heads):
        ws.write(11, c, h)
    items = items if items is not None else [
        (1, 'CSL10100\ncustomer design color box', 'Code: CSL-10100\nSize: 30cm', None, 12.5, 100, 1250, 700),
        (2, 'CSL20200', 'Material: brass', None, 8, 50, 400, 280),
        (3, 'Freight', '', None, None, None, 180, None)]
    r = 12
    for it in items:
        for c, v in enumerate(it):
            if v is not None:
                ws.write(r, c, v)
        r += 1
    ws.write(r, 0, 'Total Payment'); ws.write(r, 6, sum(i[6] for i in items))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestPiImport(AppTestCase):
    def pi_preview(self, data, name='pi.xls'):
        return self.c.j('POST', '/api/import/pi/preview', raw=data, headers={'X-Filename': name})

    def test_parse_and_customer_matching(self):
        cid = self.new_customer(company='Some Other Name', emails='BOB@acme-hotels.com')
        st, d = self.pi_preview(make_pi())
        self.assertEqual(st, 200, d)
        self.assertEqual((d['invoice_no'], d['date'], d['currency'], d['total_amount'], d['already_imported'], d['sum_matches_total']),
                         ('SL-20260101US', '2026-01-01', 'USD', 1830.0, False, True))
        self.assertEqual(d['buyer']['company'], 'ACME Hotels Ltd')
        self.assertEqual(d['matched_customer']['id'], cid)                                         # 先按邮箱（不分大小写）
        a, b, fee = d['items']
        self.assertEqual((a['sku'], a['name'], a['is_product'], a['unit_cost_cny']), ('CSL-10100', 'CSL10100', True, 7.0))   # Code 优先；rmb÷数量
        self.assertIn('customer design color box', a['spec'])
        self.assertEqual((b['sku'], b['unit_cost_cny']), ('CSL20200', 5.6))                          # 没有 Code 退回品名
        self.assertEqual((fee['is_product'], fee['amount']), (False, 180.0))
        # 按公司名精确匹配
        c2 = self.new_customer(company='Zeta Resorts')
        d2 = self.pi_preview(make_pi(buyer=('zeta resorts', 'Z', '', 'nobody@x.com', '')))[1]
        self.assertEqual(d2['matched_customer']['id'], c2)
        # 匹配不上：给候选，不自动新建
        n = self.ctx.db.scalar('SELECT COUNT(*) FROM customers')
        d3 = self.pi_preview(make_pi(buyer=('Zeta Holdings', 'Z', '', 'q@q.com', '')))[1]
        self.assertEqual((d3['matched_customer'], [x['id'] for x in d3['customer_candidates']]), (None, [c2]))
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM customers'), n)

    def test_apply_creates_everything(self):
        cid = self.new_customer(company='ACME Apply', emails='a@acme-apply.com', stage='潜在')
        d = self.pi_preview(make_pi(invoice='SL-A1US', buyer=('ACME Apply', 'A', '', 'a@acme-apply.com', '')))[1]
        d['customer_id'] = cid
        d['items'][1]['unit_cost_cny'] = 6.0                                                         # 预览里改成本
        st, r = self.c.post('/api/import/pi/apply', d)
        self.assertEqual(st, 200, r)
        self.assertEqual((r['products_new'], r['products_updated'], r['quote_no'], r['quote_total']), (2, 0, 'SL-A1US', 1830.0))
        db = self.ctx.db
        p = db.one("SELECT * FROM products WHERE sku='CSL-10100'")
        self.assertEqual((p['name'], p['cost'], p['cost_currency']), ('CSL10100', 7.0, 'CNY'))
        self.assertEqual(p['suggested_price'], round(7.0 * 0.138 * 1.25, 2))                         # 建议价仍按公式，不被 PI 成交价覆盖
        self.assertIn('Code: CSL-10100', p['spec_text'])
        self.assertEqual(db.one("SELECT cost FROM products WHERE sku='CSL20200'")['cost'], 6.0)
        q = db.one("SELECT * FROM quotes WHERE quote_no='SL-A1US'")
        self.assertEqual((q['status'], q['customer_id'], q['total'], q['created_at'][:10], q['valid_days'], q['currency']), ('accepted', cid, 1830.0, '2026-01-01', 0, 'USD'))
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id=?', (q['id'],)), 3)      # 含运费行
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id=? AND product_id IS NULL', (q['id'],)), 1)
        sell = db.one("SELECT * FROM price_history WHERE product_id=? AND price_type='sell'", (p['id'],))
        self.assertEqual((sell['price'], sell['currency'], sell['customer_id'], sell['effective_date'], sell['source']), (12.5, 'USD', cid, '2026-01-01', 'SL-A1US'))
        cost = db.one("SELECT * FROM price_history WHERE product_id=? AND price_type='cost'", (p['id'],))
        self.assertEqual((cost['price'], cost['currency'], cost['effective_date'], cost['source']), (7.0, 'CNY', '2026-01-01', 'PI导入 SL-A1US'))
        self.assertEqual(db.scalar('SELECT stage FROM customers WHERE id=?', (cid,)), '成交')
        self.assertIn('SL-A1US', db.scalar("SELECT group_concat(content) FROM notes WHERE customer_id=?", (cid,)))
        self.assertEqual(self.c.get('/api/quotes/%d' % q['id'])[1]['quote']['status'], 'accepted')
        # 同一个 PI 再导入 → 跳过
        st, r2 = self.c.post('/api/import/pi/apply', d)
        self.assertEqual((st, r2.get('skipped')), (200, True))
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM quotes WHERE quote_no='SL-A1US'"), 1)
        self.assertTrue(self.pi_preview(make_pi(invoice='SL-A1US'))[1]['already_imported'])
        # 第二张 PI：老产品不重复建，规格不覆盖，成本追加历史；客户 → 复购
        d2 = self.pi_preview(make_pi(invoice='SL-A2US', date_serial=46100, items=[(1, 'CSL10100', 'Code: CSL-10100\nChanged spec', None, 13, 20, 260, 160)]))[1]
        self.assertTrue(d2['items'][0]['exists'])
        d2['customer_id'] = cid
        r3 = self.c.post('/api/import/pi/apply', d2)[1]
        self.assertEqual((r3['products_new'], r3['products_updated']), (0, 1))
        p = db.one("SELECT * FROM products WHERE sku='CSL-10100'")
        self.assertIn('Size: 30cm', p['spec_text'])
        self.assertNotIn('Changed spec', p['spec_text'])
        self.assertEqual((p['cost'], db.scalar("SELECT COUNT(*) FROM products WHERE sku='CSL-10100'")), (8.0, 1))   # 160/20，按日期最新的成本对齐
        self.assertEqual(db.scalar('SELECT stage FROM customers WHERE id=?', (cid,)), '复购')

    def test_apply_validation(self):
        vitems = [(1, 'VAL1', 'Code: VAL-1', None, 5, 10, 50, 30), (2, 'VAL2', 'Code: VAL-2', None, 6, 10, 60, 40), (3, 'Freight', '', None, None, None, 20, None)]
        d = self.pi_preview(make_pi(invoice='SL-V1US', items=vitems))[1]
        self.assertEqual(self.c.post('/api/import/pi/apply', d)[0], 400)                              # 没指定客户
        d['customer_id'] = 987654
        self.assertEqual(self.c.post('/api/import/pi/apply', d)[0], 404)
        cid = self.new_customer(company='V Co')
        d['customer_id'] = cid
        for bad in ({'invoice_no': ''}, {'date': '2026-02-30'}, {'category_id': 99999}):
            self.assertEqual(self.c.post('/api/import/pi/apply', {**d, **bad})[0] in (400, 404), True, bad)
        items = [dict(i) for i in d['items']]
        items[0]['unit_cost_cny'] = 'abc'
        self.assertEqual(self.c.post('/api/import/pi/apply', {**d, 'items': items})[0], 400)
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM quotes WHERE quote_no='SL-V1US'"))            # 失败整体回滚
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='VAL-1'"))
        items = [dict(i) for i in d['items']]
        items[0]['include'] = False
        r = self.c.post('/api/import/pi/apply', {**d, 'items': items})[1]
        self.assertEqual(r['items'], 2)                                                                 # 勾掉的行不导入
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='VAL-1'"))

    def test_bad_files(self):
        self.assertEqual(self.pi_preview(b'x', 'a.xlsx')[0], 400)
        self.assertEqual(self.pi_preview(b'not xls at all', 'a.xls')[0], 400)
        wb = xlwt.Workbook(); wb.add_sheet('x').write(0, 0, 'hello')
        buf = io.BytesIO(); wb.save(buf)
        st, r = self.pi_preview(buf.getvalue(), 'a.xls')
        self.assertEqual(st, 400)
        self.assertIn('明细表头', r['error'])
        self.assertEqual(os.listdir(os.path.join(self.ctx.import_tmp, 'pi')), [])                      # 临时文件不残留


if __name__ == '__main__':
    unittest.main()
