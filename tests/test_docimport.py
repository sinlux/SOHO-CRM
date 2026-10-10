# -*- coding: utf-8 -*-
"""文档导入：PI（xls/xlsx）/ 供应商报价单 / 产品清单 的识别、数据清洗（查重）、确认、写入。"""
import http.client
import io
import os
import tempfile
import unittest
import urllib.parse

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XImage

from helpers import AppTestCase
from imgutil import png_bytes, PNG_RED_BOX
from test_imports import make_pi
from sinlux.imports import dedupe, docparse
from sinlux.imports import grid as gridmod

IMG_A = png_bytes(300, 300, box=(40, 40, 140, 140), box_color=(200, 30, 30))
IMG_B = png_bytes(300, 300, bg=(230, 230, 230), box=(160, 150, 280, 270), box_color=(20, 20, 200))
IMG_C = png_bytes(300, 300, bg=(250, 250, 250), box=(10, 200, 290, 280), box_color=(30, 160, 60))


def modern_pi_xlsx(images=None, invoice='BS250219SV', rows=None):
    """和旧模板不一样的 PI：标签与值在同一格、列名不同、有图片列。"""
    wb = Workbook(); ws = wb.active; ws.title = 'PI'
    ws['A1'] = 'PROFORMA INVOICE'
    ws['A3'] = 'Invoice No.: %s' % invoice
    ws['A4'] = 'Date: 2025-02-19'
    ws['A5'] = 'Buyer: Karla Distribuidora SA'
    ws['A6'] = 'Contact Person: Karla'
    ws['A7'] = 'Email: karla@karla-dist.test'
    heads = ['No.', 'Item Code', 'Description', 'Picture', 'Unit Price (USD)', 'Qty', 'Amount', 'Purchase Price (RMB)']
    for c, h in enumerate(heads, 1):
        ws.cell(row=9, column=c, value=h)
    rows = rows if rows is not None else [
        (1, 'KL-001', 'LED Downlight 10W White', None, 12.5, 100, 1250, 650),
        (2, 'KL-002', 'LED Panel Light 600x600', None, 20, 40, 800, 520),
        (3, None, 'Sea freight', None, None, None, 150, None)]
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            if v is not None:
                ws.cell(row=10 + i, column=c, value=v)
    ws.cell(row=10 + len(rows), column=1, value='TOTAL')
    ws.cell(row=10 + len(rows), column=7, value=sum(r[6] or 0 for r in rows))
    for rn, data in (images or {}).items():
        img = XImage(io.BytesIO(data)); img.anchor = 'D%d' % rn; ws.add_image(img)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def supplier_xlsx(rows, images=None, title='东莞某某灯饰有限公司 报价单'):
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
    ws['A1'] = title
    for c, h in enumerate(['货号', '图片', '品名', '规格', '单价(RMB)', 'MOQ'], 1):
        ws.cell(row=3, column=c, value=h)
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            if v is not None:
                ws.cell(row=4 + i, column=c, value=v)
    for rn, data in (images or {}).items():
        img = XImage(io.BytesIO(data)); img.anchor = 'B%d' % rn; ws.add_image(img)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


class DocBase(AppTestCase):
    def upload(self, data, name):
        return self.c.j('POST', '/api/import/doc/upload', raw=data, headers={'X-Filename': urllib.parse.quote(name)})

    def apply(self, sid, body):
        return self.c.post('/api/import/doc/%s/apply' % sid, body)

    def items_payload(self, res, over=None):
        over = over or {}
        out = []
        for it in res['items']:
            d = {k: it[k] for k in ('row', 'sku', 'name', 'spec', 'unit', 'quantity', 'unit_price', 'unit_cost_cny')}
            d['include'] = True
            d.update(over.get(it['row'], {}))
            out.append(d)
        return out


class TestDocParse(unittest.TestCase):
    def grid(self, data, name):
        d = tempfile.mkdtemp(); p = os.path.join(d, name)
        open(p, 'wb').write(data)
        return gridmod.read_grid(p, name)

    def test_classify_headers_bilingual(self):
        c = docparse.classify_header
        cases = {'Product Name': 'name', '品名': 'name', 'Description': 'name', 'Item Code': 'sku', '货号': 'sku', 'Model No.': 'sku', 'Qty': 'qty', '数量': 'qty',
                 'Unit Price (USD)': 'price', '单价(RMB)': 'price', 'Amount': 'amount', '金额': 'amount', 'Picture': 'image', '图片': 'image',
                 'RMB': 'cost_total', '采购单价': 'cost_unit', 'MOQ': 'moq', 'Specification': 'spec', '规格': 'spec', 'Unit': 'unit', 'Random column': None, '': None}
        for h, want in cases.items():
            self.assertEqual(c(h), want, h)

    def test_modern_xlsx_pi(self):
        r = docparse.parse_grid(self.grid(modern_pi_xlsx(), 'x.xlsx'))
        self.assertEqual((r['kind'], r['currency'], r['header_row']), ('pi', 'USD', 8))
        self.assertEqual((r['meta']['invoice_no'], r['meta']['date'], r['meta']['buyer'], r['meta']['email']), ('BS250219SV', '2025-02-19', 'Karla Distribuidora SA', 'karla@karla-dist.test'))
        a, b, fee = r['items']
        self.assertEqual((a['sku'], a['name'], a['quantity'], a['unit_price'], a['unit_cost_cny']), ('KL-001', 'LED Downlight 10W White', 100.0, 12.5, 6.5))
        self.assertEqual((fee['is_product'], fee['amount']), (False, 150.0))
        self.assertEqual(r['total_amount'], 2200.0)

    def test_old_xls_template_still_parses(self):
        r = docparse.parse_grid(self.grid(make_pi(), 'old.xls'))
        self.assertEqual((r['kind'], r['meta']['invoice_no'], r['meta']['date'], r['items'][0]['sku'], r['items'][0]['unit_cost_cny']), ('pi', 'SL-20260101US', '2026-01-01', 'CSL-10100', 7.0))

    def test_supplier_quote_and_unmapped_columns_go_to_spec(self):
        data = supplier_xlsx([('S-1', None, '筒灯', '10W 3000K', 8.5, 500)])
        r = docparse.parse_grid(self.grid(data, 's.xlsx'))
        self.assertEqual((r['kind'], r['currency']), ('supplier', 'CNY'))
        it = r['items'][0]
        self.assertEqual((it['sku'], it['name'], it['unit_price'], it['moq']), ('S-1', '筒灯', 8.5, 500))
        self.assertIn('10W 3000K', it['spec'])

    def test_unknown_layout_asks_for_manual_mapping(self):
        wb = Workbook(); ws = wb.active
        for r, row in enumerate([['foo', 'bar', 'baz'], ['x', 'y', 'z']], 1):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        buf = io.BytesIO(); wb.save(buf)
        r = docparse.parse_grid(self.grid(buf.getvalue(), 'u.xlsx'))
        self.assertFalse(r['ok'])
        self.assertTrue(r['warnings'])


class TestDedupeUnits(unittest.TestCase):
    def test_name_and_sku(self):
        self.assertGreaterEqual(dedupe.name_score('LED Downlight 10W White', 'LED Down light 10W white'), dedupe.HIGH_NAME)
        self.assertGreaterEqual(dedupe.name_score('LED筒灯 10W 白光', 'LED 筒灯10W 白光'), dedupe.HIGH_NAME)
        self.assertLess(dedupe.name_score('LED Downlight 10W', 'Rattan Pendant Lamp'), dedupe.MID_NAME)
        self.assertLess(dedupe.name_score('', 'x'), 0.1)
        self.assertEqual(dedupe.sku_key('CSL-10100 '), dedupe.sku_key('csl10100'))

    def test_image_hash(self):
        d = tempfile.mkdtemp()
        paths = {}
        for k, data in (('a', IMG_A), ('b', IMG_B), ('c', IMG_C)):
            paths[k] = os.path.join(d, k + '.png'); open(paths[k], 'wb').write(data)
        a2 = os.path.join(d, 'a2.png'); open(a2, 'wb').write(IMG_A)
        ha, hb, hc, ha2 = (dedupe.dhash(paths['a']), dedupe.dhash(paths['b']), dedupe.dhash(paths['c']), dedupe.dhash(a2))
        self.assertEqual(dedupe.hamming(ha, ha2), 0)
        self.assertGreater(dedupe.hamming(ha, hb), dedupe.MID_IMG)
        self.assertGreater(dedupe.hamming(ha, hc), dedupe.MID_IMG)
        white = os.path.join(d, 'w.png'); open(white, 'wb').write(png_bytes(100, 100))
        self.assertIsNone(dedupe.dhash(white))                                   # 纯白图不参与比较
        self.assertIsNone(dedupe.dhash(os.path.join(d, 'missing.png')))


class TestDocImportPI(DocBase):
    def test_xlsx_pi_upload_apply_with_images_and_margin_inputs(self):
        cid = self.new_customer(company='Karla Distribuidora SA', emails='karla@karla-dist.test')
        st, res = self.upload(modern_pi_xlsx(images={10: IMG_A, 11: IMG_B}), 'PI to Karla BS250219SV.xlsx')
        self.assertEqual(st, 200, res)
        self.assertEqual((res['kind'], res['meta']['invoice_no'], res['matched_customer']['id'], res['has_images'], res['summary']['with_image']), ('pi', 'BS250219SV', cid, True, 2))
        self.assertEqual(res['summary']['needs_decision'], 0)
        sid = res['session_id']
        st, raw, h = self.c.call('GET', '/api/import/doc/%s/image/%s' % (sid, res['items'][0]['image']))
        self.assertEqual((st, h['Content-Type'].startswith('image/')), (200, True))
        self.assertEqual(self.c.call('GET', '/api/import/doc/%s/image/..%%2fupload.xlsx' % sid)[0], 404)
        st, r = self.apply(sid, {'customer_id': cid, 'items': self.items_payload(res)})
        self.assertEqual(st, 200, r)
        self.assertEqual((r['products_new'], r['quote_no'], r['quote_total']), (2, 'BS250219SV', 2200.0))
        db = self.ctx.db
        p = db.one("SELECT * FROM products WHERE sku='KL-001'")
        self.assertEqual((p['name'], p['cost'], p['cost_currency']), ('LED Downlight 10W White', 6.5, 'CNY'))
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM product_images WHERE product_id=? AND normalized=1', (p['id'],)), 1)        # PI 里的图片进了产品相册
        q = db.one("SELECT * FROM quotes WHERE quote_no='BS250219SV'")
        self.assertEqual((q['status'], q['created_at'][:10], q['total']), ('accepted', '2025-02-19', 2200.0))
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id=?', (q['id'],)), 3)
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM price_history WHERE source='BS250219SV' AND price_type='sell'"), 2)
        self.assertEqual(self.c.get('/api/import/doc/%s/image/x.jpg' % sid)[0], 404)                                               # 会话用完即清

    def test_old_xls_via_new_endpoint_and_products_only(self):
        cid = self.new_customer(company='ACME Docs', emails='bob@acme-hotels.com')
        st, res = self.upload(make_pi(invoice='SL-D1US'), 'old.xls')
        self.assertEqual((st, res['kind'], res['is_xls'], res['matched_customer']['id']), (200, 'pi', True, cid))
        st, r = self.apply(res['session_id'], {'customer_id': cid, 'create_quote': False, 'items': self.items_payload(res)})
        self.assertEqual((st, r['products_new']), (200, 2))
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM quotes WHERE quote_no='SL-D1US'"))                                         # 只导入产品，不生成成交单
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM price_history WHERE source LIKE '%old.xls%' AND price_type='cost'"), 2)

    def test_validation_and_duplicate_invoice(self):
        cid = self.new_customer(company='Val Co')
        st, res = self.upload(modern_pi_xlsx(invoice='DUP-1', rows=[(1, 'V-1', 'Valve lamp', None, 5, 10, 50, 20)]), 'a.xlsx')
        sid = res['session_id']
        self.assertEqual(self.apply(sid, {'items': self.items_payload(res)})[0], 400)                                              # 没指定客户
        self.assertEqual(self.apply(sid, {'customer_id': 987654, 'items': self.items_payload(res)})[0], 404)
        self.assertEqual(self.apply(sid, {'customer_id': cid, 'currency': 'EUR', 'items': self.items_payload(res)})[0], 400)
        self.assertEqual(self.apply(sid, {'customer_id': cid, 'items': [{'row': 999}]})[0], 400)
        self.assertEqual(self.apply(sid, {'customer_id': cid, 'items': self.items_payload(res, {10: {'unit_price': 'abc'}})})[0], 400)
        self.assertEqual(self.apply(sid, {'customer_id': cid, 'items': self.items_payload(res, {10: {'quantity': 0}})})[0], 400)
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='V-1'"))                                                 # 校验失败一条都不写
        self.assertEqual(self.apply(sid, {'customer_id': cid, 'items': self.items_payload(res)})[0], 200)
        st, res2 = self.upload(modern_pi_xlsx(invoice='DUP-1', rows=[(1, 'V-1', 'Valve lamp', None, 5, 10, 50, 20)]), 'a.xlsx')
        self.assertTrue(res2['already_imported'])
        st, r = self.apply(res2['session_id'], {'customer_id': cid, 'items': self.items_payload(res2)})
        self.assertEqual((st, r.get('skipped')), (200, True))

    def test_wrong_extension_gives_clear_error_even_for_big_upload(self):
        port = self.port
        conn = http.client.HTTPConnection('127.0.0.1', port)
        conn.request('POST', '/api/import/doc/upload', body=b'x' * (25 * 1024 * 1024), headers={'X-Filename': 'a.pdf'})
        r = conn.getresponse()
        body = r.read()
        self.assertEqual(r.status, 400)                                                                                            # 不是「Failed to fetch」
        self.assertIn('xls', body.decode('utf-8'))
        conn.request('POST', '/api/import/doc/upload', body=b'not excel', headers={'X-Filename': 'a.xlsx'})
        r = conn.getresponse()
        self.assertEqual(r.status, 400)
        r.read()
        conn.close()


class TestDocImportSupplier(DocBase):
    def test_supplier_quote_creates_products_vendor_quotes_and_images(self):
        data = supplier_xlsx([('SQ-1', None, '筒灯 10W', '3000K', 8.5, 500), ('SQ-2', None, '射灯 7W', '4000K', 6.2, 300)], images={4: IMG_A, 5: IMG_C})
        st, res = self.upload(data, '某某灯饰 报价.xlsx')
        self.assertEqual((st, res['kind'], res['currency'], res['summary']['with_image']), (200, 'supplier', 'CNY', 2))
        sid = res['session_id']
        self.assertEqual(self.apply(sid, {'items': self.items_payload(res)})[0], 400)                                               # 没填供应商
        st, r = self.apply(sid, {'supplier_name': '某某灯饰厂', 'project': 'Hotel Azul', 'items': self.items_payload(res)})
        self.assertEqual((st, r['products_new'], r['supplier_quotes']), (200, 2, 2), r)
        db = self.ctx.db
        v = db.one("SELECT * FROM suppliers WHERE name='某某灯饰厂'")
        self.assertIsNotNone(v)
        p = db.one("SELECT * FROM products WHERE sku='SQ-1'")
        self.assertEqual((p['cost'], p['cost_currency'], p['supplier'], p['moq']), (8.5, 'CNY', '某某灯饰厂', 500))
        q = db.one('SELECT * FROM supplier_quotes WHERE product_id=?', (p['id'],))
        self.assertEqual((q['price_cny'], q['project'], q['supplier_id']), (8.5, 'Hotel Azul', v['id']))
        self.assertEqual(db.scalar('SELECT COUNT(*) FROM product_images WHERE product_id=?', (p['id'],)), 1)
        # 再导一次：同 SKU 自动并入（不新建、不覆盖成本），供应商报价追加
        st, res2 = self.upload(data, 'again.xlsx')
        self.assertEqual(res2['items'][0]['suggest']['action'], 'merge')
        st, r2 = self.apply(res2['session_id'], {'supplier_name': '某某灯饰厂', 'items': self.items_payload(res2)})
        self.assertEqual((r2['products_new'], r2['products_merged']), (0, 2))
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM products WHERE sku='SQ-1'"), 1)
        self.assertEqual(db.one("SELECT cost FROM products WHERE sku='SQ-1'")['cost'], 8.5)

    def test_usd_supplier_quote_is_refused_clearly(self):
        data = supplier_xlsx([('U-1', None, 'x', 'y', 5, 10)])
        st, res = self.upload(data, 'u.xlsx')
        st, r = self.c.post('/api/import/doc/%s/reparse' % res['session_id'], {'kind': 'supplier', 'currency': 'USD'})
        self.assertEqual((st, r['currency']), (200, 'USD'))
        st, r = self.apply(res['session_id'], {'kind': 'supplier', 'currency': 'USD', 'supplier_name': 'X', 'items': self.items_payload(r)})
        self.assertEqual(st, 400)
        self.assertIn('人民币', r['error'])


class TestDocImportCleaning(DocBase):
    def test_similar_products_must_be_decided(self):
        old = self.new_product(sku='OLD-DL', name='LED Downlight 10W White', category_id=self.cat_id('lighting'))
        up = self.c.j('POST', '/api/images/upload', raw=IMG_B)[1]['image']
        self.c.put('/api/products/%d' % old, {'images': [up['id']]})
        data = supplier_xlsx([('N-1', None, 'LED Down light 10W white', 'x', 5, 100),        # 名称高度相似
                              ('N-2', None, 'Totally different thing', 'x', 6, 100),         # 图片相同但名称不同
                              ('N-3', None, 'Another unrelated item', 'x', 7, 100)],         # 无关
                             images={5: IMG_B})
        st, res = self.upload(data, 'sim.xlsx')
        by = {it['sku']: it for it in res['items']}
        self.assertEqual(by['N-1']['suggest']['action'], 'review')
        self.assertIn('name_high', by['N-1']['matches'][0]['reasons'])
        self.assertEqual(by['N-2']['suggest']['action'], 'review')
        self.assertIn('image_high', by['N-2']['matches'][0]['reasons'])
        self.assertIsNone(by['N-3']['suggest'])
        self.assertEqual(res['summary']['needs_decision'], 2)
        sid = res['session_id']
        body = {'supplier_name': '查重供应商', 'items': self.items_payload(res)}
        st, r = self.apply(sid, body)
        self.assertEqual(st, 409)                                                                                                  # 没决定就不让导入
        self.assertIn('2 个疑似重复', r['error'])
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='N-3'"))
        decide = {by['N-1']['row']: {'decision': {'action': 'merge', 'target_id': old}}, by['N-2']['row']: {'decision': {'action': 'skip'}}}
        st, r = self.apply(sid, {**body, 'items': self.items_payload(res, decide)})
        self.assertEqual((st, r['products_new'], r['products_merged'], r['products_skipped']), (200, 1, 1, 1), r)
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM products WHERE sku IN ('N-1','N-2')"), 0)
        self.assertIsNotNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='N-3'"))

    def test_new_with_conflicting_sku_and_in_file_duplicates(self):
        self.new_product(sku='TAKEN-1', name='Existing taken')
        data = supplier_xlsx([('TAKEN-1', None, 'Brand new name', 'x', 5, 10), ('DUP-A', None, 'Same thing', 'x', 5, 10), ('DUP-A', None, 'Same thing', 'x', 5.5, 10)])
        st, res = self.upload(data, 'c.xlsx')
        by = [it for it in res['items']]
        self.assertEqual(by[0]['suggest']['action'], 'merge')                                       # 同 SKU：默认并入已有产品
        dup = by[2]
        self.assertEqual(dup['suggest']['action'], 'link_row')
        self.assertIn('price_conflict', [f['code'] for f in dup['flags']])
        sid = res['session_id']
        pre = {by[0]['row']: {'decision': {'action': 'new'}}, dup['row']: {'decision': {'action': 'link_row', 'target_row': by[1]['row']}}}
        body = {'supplier_name': 'S', 'items': self.items_payload(res, pre)}
        st, r = self.apply(sid, body)
        self.assertEqual(st, 400)
        self.assertIn('TAKEN-1', r['error'])                                                          # 想「新建」却和已有 SKU 冲突
        self.assertIsNone(self.ctx.db.one("SELECT 1 FROM products WHERE sku='DUP-A'"))
        decide = {dup['row']: {'decision': {'action': 'link_row', 'target_row': by[1]['row']}}}
        st, r = self.apply(sid, {'supplier_name': 'S', 'items': self.items_payload(res, decide)})
        self.assertEqual((st, r['products_new'], r['products_merged']), (200, 1, 1), r)
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM products WHERE sku='DUP-A'"), 1)

    def test_manual_mapping_fallback_and_error_flags(self):
        wb = Workbook(); ws = wb.active
        rows = [['', '', ''], ['代号', '叫什么', '多少钱', '几个'], ['M-1', '手动识别灯', 9.9, 10]]
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
        buf = io.BytesIO(); wb.save(buf)
        st, res = self.upload(buf.getvalue(), 'odd.xlsx')
        self.assertFalse(res['ok'])
        self.assertTrue(res['grid_preview'])
        st, r = self.c.post('/api/import/doc/%s/reparse' % res['session_id'], {'header_row': 2, 'kind': 'list', 'currency': 'CNY', 'columns': {'sku': 0, 'name': 1, 'price': 2, 'qty': 3}})
        self.assertEqual((st, r['ok'], r['items'][0]['sku'], r['items'][0]['unit_price'], r['kind']), (200, True, 'M-1', 9.9, 'list'))
        st, r2 = self.apply(res['session_id'], {'kind': 'list', 'currency': 'CNY', 'items': self.items_payload(r)})
        self.assertEqual((st, r2['products_new']), (200, 1))
        p = self.ctx.db.one("SELECT cost, cost_currency FROM products WHERE sku='M-1'")
        self.assertEqual((p['cost'], p['cost_currency']), (9.9, 'CNY'))

    def test_flags_below_cost_and_missing_fields(self):
        cid = self.new_customer(company='Flag Co')
        data = modern_pi_xlsx(invoice='FLG-1', rows=[(1, 'F-1', 'Loss maker', None, 1.0, 10, 10, 500), (2, None, 'No price no sku', None, None, 5, 5, None)])
        st, res = self.upload(data, 'f.xlsx')
        self.assertIn('below_cost', {f['code'] for f in res['items'][0]['flags']})
        codes2 = {f['code'] for f in res['items'][1]['flags']}
        self.assertTrue({'price_missing'} <= codes2, codes2)
        self.assertTrue(res['summary']['errors'] >= 1 and res['summary']['warnings'] >= 1)


class TestLibraryScan(DocBase):
    def test_scan_finds_name_sku_and_image_groups(self):
        a = self.new_product(sku='SCAN-A1', name='Rattan Pendant Lamp 45cm')
        b = self.new_product(sku='SCAN-A2', name='Rattan pendant lamp 45 cm')          # 名称几乎一样
        c = self.new_product(sku='SCAN-C1', name='Walnut dining chair')
        d = self.new_product(sku='scan c1', name='Something else entirely')               # SKU 近似
        e = self.new_product(sku='SCAN-E1', name='Bamboo basket')
        f = self.new_product(sku='SCAN-F1', name='Totally different name')                # 图片相同
        for pid in (e, f):
            up = self.c.j('POST', '/api/images/upload', raw=IMG_C)[1]['image']
            self.c.put('/api/products/%d' % pid, {'images': [up['id']]})
        g = self.new_product(sku='SCAN-G1', name='Unique ottoman')
        r = self.c.get('/api/products/duplicates')[1]
        groups = [sorted(m['id'] for m in grp['members']) for grp in r['groups']]
        self.assertIn(sorted([a, b]), groups)
        self.assertIn(sorted([c, d]), groups)
        self.assertIn(sorted([e, f]), groups)
        self.assertFalse(any(g in grp for grp in groups))
        reasons = {tuple(sorted(m['id'] for m in grp['members'])): grp['reasons'] for grp in r['groups']}
        self.assertIn('image_high', reasons[tuple(sorted([e, f]))])
        # 合并走已有的合并接口：确认后才合并
        st, m = self.c.post('/api/products/merge', {'survivor_id': a, 'merge_ids': [b]})
        self.assertEqual(st, 200, m)
        groups = [sorted(x['id'] for x in grp['members']) for grp in self.c.get('/api/products/duplicates')[1]['groups']]
        self.assertNotIn(sorted([a, b]), groups)


if __name__ == '__main__':
    unittest.main()
