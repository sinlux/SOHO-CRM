# -*- coding: utf-8 -*-
"""客户 Excel 导入（先预览后写入、可恢复会话、原子写入）与导出。"""
import io
import os
import unittest
import urllib.parse

from openpyxl import Workbook, load_workbook

from helpers import AppTestCase
from sinlux.customers.xlsx_io import CustomerImporter

HEADER = ['LV', 'Country', 'Name', 'Company', 'Website', 'Email', 'Main business', 'Address', 'Whatsapp',
          'Linkedin', 'Facebook', 'Remark']


def make_xlsx(rows, header=HEADER, extra_sheet=False):
    wb = Workbook()
    ws = wb.active
    ws.title = '总表'
    ws.append(header)
    for r in rows:
        ws.append(r)
    if extra_sheet:
        wb.create_sheet('第二张').append(['junk'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestCustomerImport(AppTestCase):
    def upload(self, data, name='客户汇总.xlsx', remap=False):
        h = {'X-Filename': urllib.parse.quote(name)}
        return self.c.j('POST', '/api/import/customers/upload?remap_lv=%d' % remap, raw=data, headers=h)

    def count(self, where='1=1'):
        return self.ctx.db.scalar('SELECT COUNT(*) FROM customers WHERE ' + where)

    def test_preview_writes_nothing_then_apply_writes(self):
        rows = [[5, 'Germany', 'Hans', 'Imp Co A', 'www.a.com; a.com', 'h@a.com, h@a.com sales@a.com', 'LED', 'Berlin', '+49 1', '', '', '长期客户'],
                [3, 'USA', 'Bob', 'Imp Co B', 'www.b.com', 'bob@b.com', '', '', '', '', '', '']]
        before = self.count()
        st, p = self.upload(make_xlsx(rows))
        self.assertEqual(st, 200, p)
        self.assertEqual((p['to_import'], p['duplicates'], p['skipped'], p['sheet']), (2, 0, 0, '总表'))
        self.assertEqual(self.count(), before)                         # 预览阶段不写库
        self.assertEqual(p['sample'][0]['data']['emails'], 'h@a.com sales@a.com')       # 多邮箱拆分去重
        self.assertEqual(p['sample'][0]['data']['website'], 'www.a.com a.com')
        st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id']})
        self.assertEqual((st, r['imported']), (200, 2), r)
        self.assertEqual(self.count(), before + 2)
        cid = self.ctx.db.one("SELECT id FROM customers WHERE company='Imp Co A'")['id']
        d = self.c.get('/api/customers/%d' % cid)[1]
        self.assertEqual(d['customer']['lv'], 5)
        self.assertEqual([n['content'] for n in d['notes']], ['【导入时的Remark】长期客户'])  # Remark 变备注
        # 会话用完即清理，不能重复 apply
        st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id']})
        self.assertEqual(st, 404)

    def test_duplicates_against_db_and_within_file(self):
        self.new_customer(company='Existing Co', emails='old@existing.com')
        rows = [[1, '', '', 'existing co', '', '', '', '', '', '', '', ''],             # 公司名重复(大小写不敏感)
                [1, '', '', 'Another', '', 'OLD@existing.com', '', '', '', '', '', ''],   # 邮箱重复
                [1, '', '', 'Fresh Co', '', 'f@fresh.com', '', '', '', '', '', ''],
                [1, '', '', 'FRESH CO', '', '', '', '', '', '', '', ''],                 # 文件内重复
                [1, '', '', 'Fresh Two', '', 'F@fresh.com', '', '', '', '', '', '']]     # 文件内邮箱重复
        st, p = self.upload(make_xlsx(rows))
        self.assertEqual((p['to_import'], p['duplicates']), (1, 4))
        self.assertEqual(len([x for x in p['problems'] if '重复' in x]), 4)
        self.assertIn('第2行', p['problems'][0])                        # 行号按 Excel 实际行号(表头为第1行)

    def test_lv_validation_and_remap(self):
        rows = [[1, '', '', 'Lv One', '', '', '', '', '', '', '', ''], [6, '', '', 'Lv Six', '', '', '', '', '', '', '', ''],
                ['abc', '', '', 'Lv Bad', '', '', '', '', '', '', '', ''], [9, '', '', 'Lv Nine', '', '', '', '', '', '', '', ''],
                [2.5, '', '', 'Lv Frac', '', '', '', '', '', '', '', ''], [None, '', '', 'Lv None', '', '', '', '', '', '', '', '']]
        st, p = self.upload(make_xlsx(rows))
        lvs = {r['data']['company']: r['data']['lv'] for r in p['sample']}
        self.assertEqual(lvs, {'Lv One': 1, 'Lv Six': 6, 'Lv Bad': None, 'Lv Nine': None, 'Lv Frac': None, 'Lv None': None})
        self.assertEqual(len([x for x in p['problems'] if 'LV' in x]), 3)         # 越界/非数字/小数都要报告
        st, p2 = self.c.post('/api/import/customers/preview', {'session_id': p['session_id'], 'remap_lv': True})
        lvs = {r['data']['company']: r['data']['lv'] for r in p2['sample']}
        self.assertEqual((lvs['Lv One'], lvs['Lv Six']), (6, 1))                  # 1↔6 互换
        st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id'], 'remap_lv': True})
        self.assertEqual(self.ctx.db.one("SELECT lv FROM customers WHERE company='Lv One'")['lv'], 6)

    def test_header_variants_blank_rows_and_skips(self):
        header = ['company', 'EMAIL', 'main business', 'lv', 'junk col']
        rows = [['Case Co', 'c@case.com', 'x', 4, 'ignored'], [None, None, None, None, None], ['', '  ', '', '', ''],
                [None, None, 'only business', 3, '']]
        st, p = self.upload(make_xlsx(rows, header))
        self.assertEqual((st, p['to_import'], p['skipped']), (200, 1, 1))
        self.assertTrue(any('全为空' in x for x in p['problems']))

    def test_unrecognized_or_bad_files_rejected(self):
        root = os.path.join(self.ctx.import_tmp, 'customers')
        sessions = lambda: set(os.listdir(root)) if os.path.isdir(root) else set()
        before = sessions()
        st, r = self.upload(make_xlsx([['x']], header=['foo', 'bar']))
        self.assertEqual(st, 400)
        self.assertIn('表头无法识别', r['error'])
        st, r = self.upload(b'this is not an excel file')
        self.assertEqual(st, 400)
        st, r = self.upload(b'x', name='data.csv')
        self.assertEqual(st, 400)
        st, r = self.upload(b'')
        self.assertEqual(st, 400)
        self.assertEqual(sessions(), before)                  # 失败的上传不留垃圾

    def test_session_id_validation(self):
        for sid in ('../../etc', '..', 'zzzzzzzzzzzz', '', 'abcdef123456'):
            st, r = self.c.post('/api/import/customers/apply', {'session_id': sid})
            self.assertIn(st, (400, 404), sid)

    def test_session_survives_restart(self):
        """旧版会话在内存里，重启即丢；新版落盘，换一个导入器实例（模拟重启）仍可继续。"""
        st, p = self.upload(make_xlsx([[2, 'FR', '', 'Restart Co', '', 'r@restart.com', '', '', '', '', '', '']]))
        fresh = CustomerImporter(self.ctx.db, self.ctx.customers, self.ctx.importer.tmp_root)
        res = fresh.apply(p['session_id'])
        self.assertEqual(res['imported'], 1)
        self.assertEqual(self.count("company='Restart Co'"), 1)

    def test_apply_is_atomic(self):
        st, p = self.upload(make_xlsx([[1, '', '', 'Atomic %d' % i, '', '', '', '', '', '', '', ''] for i in range(5)]))
        orig, n = self.ctx.customers.create, [0]

        def flaky(d):
            n[0] += 1
            if n[0] == 4:
                raise RuntimeError('disk full')
            return orig(d)
        self.ctx.customers.create = flaky
        try:
            st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id']})
        finally:
            self.ctx.customers.create = orig
        self.assertEqual(st, 500)
        self.assertEqual(self.count("company LIKE 'Atomic %'"), 0)       # 全部回滚，不会导入一半
        st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id']})
        self.assertEqual((st, r['imported']), (200, 5))                  # 会话保留，修好后可重试

    def test_cleanup_old_sessions(self):
        st, p = self.upload(make_xlsx([[1, '', '', 'Old Session', '', '', '', '', '', '', '', '']]))
        d = os.path.join(self.ctx.importer.tmp_root, p['session_id'])
        old = os.path.getmtime(d) - 48 * 3600
        os.utime(d, (old, old))
        self.ctx.importer.cleanup_old(24)
        self.assertFalse(os.path.exists(d))

    def test_large_file_streams(self):
        rows = [[(i % 6) + 1, 'C', 'N%d' % i, 'Bulk Co %d' % i, 'www.bulk%d.com' % i, 'b%d@bulk.com' % i, 'x', '', '', '', '', '']
                for i in range(3000)]
        st, p = self.upload(make_xlsx(rows))
        self.assertEqual((st, p['to_import']), (200, 3000))
        st, r = self.c.post('/api/import/customers/apply', {'session_id': p['session_id']})
        self.assertEqual(r['imported'], 3000)


class TestExport(AppTestCase):
    def test_export_roundtrip_and_formula_safety(self):
        a = self.new_customer(company='Exp A', lv=6, stage='成交', emails='a@exp.com', whatsapp='+86 138 0000', country='CN')
        self.new_customer(company='=HYPERLINK("http://evil","x")', lv=2, main_business='+SUM(1)')
        st, r = self.c.post('/api/customers/export')
        self.assertEqual(st, 200, r)
        st, raw, h = self.c.call('GET', r['url'])
        self.assertEqual(st, 200)
        self.assertIn('spreadsheetml', h['Content-Type'])
        self.assertIn('attachment', h['Content-Disposition'])
        ws = load_workbook(io.BytesIO(raw)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][:5], ('LV', '阶段', '国家', '联系人', '公司'))
        self.assertEqual(len(rows) - 1, self.ctx.db.scalar('SELECT COUNT(*) FROM customers'))
        self.assertEqual(rows[1][0], 6)                                    # LV 高的在前
        by_company = {r[4]: r for r in rows[1:]}
        self.assertEqual(by_company['Exp A'][7], '+86 138 0000')           # 电话 + 号不被改写
        evil = [r for r in rows[1:] if r[4] and 'HYPERLINK' in r[4]][0]
        cell = [c for row in ws.iter_rows() for c in row if c.value and 'HYPERLINK' in str(c.value)][0]
        self.assertEqual(cell.data_type, 's')                              # 当文本存，不是公式
        self.assertEqual(evil[11], '+SUM(1)')

    def test_export_download_validation(self):
        for p in ('/exports/..%2fcrm.db', '/exports/x.pdf', '/exports/missing.xlsx'):
            self.assertEqual(self.c.call('GET', p)[0], 404, p)


if __name__ == '__main__':
    unittest.main()
