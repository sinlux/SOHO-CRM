# -*- coding: utf-8 -*-
"""第三批：报价单（创建/编辑/状态联动/价格历史/阶段推进）+ PDF / Excel / WhatsApp 输出 + 抬头设置。"""
import io
import os
import re
import subprocess
import tempfile
import time
import unittest
import urllib.parse

from openpyxl import load_workbook

from helpers import AppTestCase
from imgutil import png_bytes
from legacy_db import build_legacy_db
from sinlux.quotes.settings import render_template, DEFAULT_WA_TEMPLATE

TODAY = time.strftime('%Y%m%d')


def pdf_text(raw):
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
        f.write(raw)
    try:
        return subprocess.run(['pdftotext', '-layout', f.name, '-'], capture_output=True, text=True, check=True).stdout
    finally:
        os.remove(f.name)


def pdf_pages(raw):
    from pypdf import PdfReader
    return len(PdfReader(io.BytesIO(raw)).pages)


class QuoteBase(AppTestCase):
    def mk_quote(self, cid=None, items=None, **kw):
        cid = cid or self.new_customer(company='Quote Buyer %d' % self.ctx.db.scalar('SELECT COUNT(*) FROM customers'))
        body = {'customer_id': cid, 'currency': 'USD', 'items': items or [{'sku': 'A-1', 'name': 'Lamp', 'quantity': 10, 'unit_price': 2.5}]}
        body.update(kw)
        st, r = self.c.post('/api/quotes', body)
        self.assertEqual(st, 200, r)
        return r['id'], cid

    def quote(self, qid):
        st, r = self.c.get('/api/quotes/%d' % qid)
        self.assertEqual(st, 200, r)
        return r['quote']

    def stage(self, cid):
        return self.ctx.db.one('SELECT stage FROM customers WHERE id=?', (cid,))['stage']

    def sells(self, pid=None):
        sql = "SELECT * FROM price_history WHERE price_type='sell'" + (' AND product_id=?' if pid else '') + ' ORDER BY id'
        return self.ctx.db.query(sql, (pid,) if pid else ())


class TestQuoteCrud(QuoteBase):
    def test_create_numbering_snapshot_and_timeline(self):
        pid = self.new_product(sku='SNAP-1', name='Original Name', cost=10, cost_currency='USD', spec_text='orig spec')
        cid = self.new_customer(company='Snapshot Buyer')
        qid, _ = self.mk_quote(cid, items=[{'product_id': pid, 'sku': 'SNAP-1', 'name': 'Original Name', 'spec': 'orig spec', 'quantity': 3, 'unit': 'set', 'unit_price': 12.5, 'remark': 'r1'},
                                            {'sku': '', 'name': 'Custom item not in library', 'quantity': 1, 'unit_price': 100}])
        q = self.quote(qid)
        self.assertRegex(q['quote_no'], r'^SLQ-%s-\d{3}$' % TODAY)
        self.assertEqual((q['status'], q['status_label'], q['currency'], q['valid_days'], q['total']), ('sent', '已发送', 'USD', 30, 137.5))
        self.assertEqual([(i['sku'], i['quantity'], i['unit'], i['amount']) for i in q['items']], [('SNAP-1', 3.0, 'set', 37.5), ('', 1.0, 'pcs', 100.0)])
        self.assertEqual(q['items'][1]['product_id'], None)
        self.assertEqual((q['company'], q['valid_until'] > q['created_at'][:10]), ('Snapshot Buyer', True))
        # 快照：之后改产品，历史报价不变
        self.c.put('/api/products/%d' % pid, {'name': 'Renamed', 'spec_text': 'new spec', 'cost': 99})
        again = self.quote(qid)
        self.assertEqual((again['items'][0]['name'], again['items'][0]['spec'], again['items'][0]['unit_price']), ('Original Name', 'orig spec', 12.5))
        notes = self.c.get('/api/customers/%d' % cid)[1]['notes']
        self.assertTrue(any(q['quote_no'] in n['content'] and '共2项' in n['content'] and 'SNAP-1' in n['content'] for n in notes))
        # 编号递增
        q2, _ = self.mk_quote(cid)
        self.assertEqual(int(self.quote(q2)['quote_no'][-3:]), int(q['quote_no'][-3:]) + 1)

    def test_numbering_uses_numeric_max_beyond_999(self):
        cid = self.new_customer()
        self.ctx.db.execute("INSERT INTO quotes(quote_no,customer_id) VALUES(?,?)", ('SLQ-%s-999' % TODAY, cid))
        qid, _ = self.mk_quote(cid)
        self.assertEqual(self.quote(qid)['quote_no'], 'SLQ-%s-1000' % TODAY)
        qid2, _ = self.mk_quote(cid)
        self.assertEqual(self.quote(qid2)['quote_no'], 'SLQ-%s-1001' % TODAY)             # 字符串排序会把 1000 排在 999 前面，必须按数值取最大

    def test_money_rounding_total_equals_sum_of_displayed_lines(self):
        qid, _ = self.mk_quote(items=[{'name': 'a', 'quantity': 3, 'unit_price': 0.335}, {'name': 'b', 'quantity': 3, 'unit_price': 0.335},
                                      {'name': 'c', 'quantity': 7, 'unit_price': 1.005}, {'name': 'd', 'quantity': 0.5, 'unit_price': 0.01}])
        q = self.quote(qid)
        self.assertEqual([i['amount'] for i in q['items']], [1.01, 1.01, 7.04, 0.01])    # 四舍五入到分（ROUND_HALF_UP，不是银行家舍入）
        self.assertEqual(q['total'], 9.07)
        self.assertEqual(round(sum(i['amount'] for i in q['items']), 2), q['total'])

    def test_validation(self):
        cid = self.new_customer()
        ok = {'name': 'x', 'quantity': 1, 'unit_price': 1}
        bad = [({'items': [ok]}, '请选择客户'), ({'customer_id': 987654, 'items': [ok]}, '客户不存在'), ({'customer_id': cid, 'items': []}, '至少要有一个'),
               ({'customer_id': cid, 'items': [{}]}, '至少要有一个'), ({'customer_id': cid, 'currency': 'GBP', 'items': [ok]}, '币种'),
               ({'customer_id': cid, 'items': [{**ok, 'quantity': 0}]}, '数量'), ({'customer_id': cid, 'items': [{**ok, 'quantity': -2}]}, '数量'),
               ({'customer_id': cid, 'items': [{**ok, 'quantity': 'abc'}]}, '数量'), ({'customer_id': cid, 'items': [{**ok, 'unit_price': -1}]}, '单价'),
               ({'customer_id': cid, 'items': [{**ok, 'unit_price': 'nan'}]}, '单价'), ({'customer_id': cid, 'items': [{**ok, 'unit_price': 1e15}]}, '过大'),
               ({'customer_id': cid, 'items': [{'unit_price': 5, 'quantity': 1}]}, '没有产品名称'), ({'customer_id': cid, 'items': [{**ok, 'product_id': 987654}]}, '产品不存在'),
               ({'customer_id': cid, 'valid_days': -1, 'items': [ok]}, '有效期'), ({'customer_id': cid, 'valid_days': 99999, 'items': [ok]}, '有效期'),
               ({'customer_id': cid, 'status': 'weird', 'items': [ok]}, '状态'), ({'customer_id': cid, 'items': [ok] * 201}, '最多 200')]
        for body, kw in bad:
            st, r = self.c.post('/api/quotes', body)
            self.assertEqual(st, 400 if 'customer_id' not in body or kw != '客户不存在' else 404, (kw, r))
            self.assertIn(kw, r['error'])
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM quotes WHERE customer_id=?', (cid,)), 0)    # 任何一次失败都不留半成品

    def test_blank_rows_ignored_and_product_fills_blank_name(self):
        pid = self.new_product(sku='FILL-1', name='Filled Name')
        qid, _ = self.mk_quote(items=[{'product_id': pid, 'quantity': 2, 'unit_price': 4}, {'sku': '', 'name': '', 'quantity': 1, 'unit_price': 0}, {}])
        q = self.quote(qid)
        self.assertEqual([(i['sku'], i['name']) for i in q['items']], [('FILL-1', 'Filled Name')])

    def test_atomic_creation_rolls_back_everything(self):
        cid = self.new_customer()
        orig = self.ctx.customers.add_system_note
        self.ctx.customers.add_system_note = lambda *a, **k: 1 / 0
        try:
            st, r = self.c.post('/api/quotes', {'customer_id': cid, 'items': [{'name': 'x', 'quantity': 1, 'unit_price': 1}]})
        finally:
            self.ctx.customers.add_system_note = orig
        self.assertEqual(st, 500)
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM quotes WHERE customer_id=?', (cid,)), 0)
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id NOT IN (SELECT id FROM quotes)'), 0)

    def test_update_replaces_items_keeps_number_and_blocks_when_accepted(self):
        qid, cid = self.mk_quote(items=[{'name': 'old', 'quantity': 1, 'unit_price': 1}], lead_time='10 days')
        no = self.quote(qid)['quote_no']
        st, r = self.c.put('/api/quotes/%d' % qid, {'currency': 'EUR', 'lead_time': '20 days', 'valid_days': 15,
                                                    'items': [{'name': 'new1', 'quantity': 2, 'unit_price': 5}, {'name': 'new2', 'quantity': 1, 'unit_price': 3}]})
        self.assertEqual(st, 200, r)
        q = r['quote']
        self.assertEqual((q['quote_no'], q['currency'], q['lead_time'], q['valid_days'], q['total'], [i['name'] for i in q['items']]), (no, 'EUR', '20 days', 15, 13.0, ['new1', 'new2']))
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id=?', (qid,)), 2)
        self.assertTrue(any('已修改' in n['content'] for n in self.c.get('/api/customers/%d' % cid)[1]['notes']))
        other = self.new_customer()
        self.assertEqual(self.c.put('/api/quotes/%d' % qid, {'customer_id': other, 'items': [{'name': 'x', 'quantity': 1, 'unit_price': 1}]})[0], 400)
        self.assertEqual(self.c.put('/api/quotes/%d' % qid, {'items': []})[0], 400)
        self.assertEqual(self.quote(qid)['total'], 13.0)                                  # 校验失败不改动原单
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'accepted'})
        st, r = self.c.put('/api/quotes/%d' % qid, {'items': [{'name': 'z', 'quantity': 1, 'unit_price': 1}]})
        self.assertEqual(st, 409)
        self.assertIn('已成交', r['error'])
        self.assertEqual(self.c.put('/api/quotes/987654', {'items': [{'name': 'z'}]})[0], 404)

    def test_duplicate_makes_independent_draft(self):
        pid = self.new_product(sku='DQ-1')
        qid, cid = self.mk_quote(items=[{'product_id': pid, 'sku': 'DQ-1', 'name': 'DQ', 'quantity': 4, 'unit_price': 7}], payment_terms='T/T', notes='n1')
        r = self.c.post('/api/quotes/%d/duplicate' % qid)[1]
        d = self.quote(r['id'])
        self.assertEqual((d['status'], d['payment_terms'], d['notes'], d['total'], d['items'][0]['product_id']), ('draft', 'T/T', 'n1', 28.0, pid))
        self.assertNotEqual(d['quote_no'], self.quote(qid)['quote_no'])
        self.c.delete('/api/products/%d' % pid, {'confirm': True})          # 原产品被删后仍可复制（明细是快照）
        r2 = self.c.post('/api/quotes/%d/duplicate' % qid)
        self.assertEqual(r2[0], 200, r2)
        self.assertIsNone(self.quote(r2[1]['id'])['items'][0]['product_id'])


class TestQuoteListSearch(QuoteBase):
    def test_filters_search_pagination_and_flags(self):
        a = self.new_customer(company='Alpha Imports', name='Ann')
        b = self.new_customer(company='Beta Trading', name='Bob')
        q1, _ = self.mk_quote(a, items=[{'sku': 'ZZ-OAK', 'name': 'Oak table', 'quantity': 1, 'unit_price': 1}])
        q2, _ = self.mk_quote(b, items=[{'sku': 'YY', 'name': 'Glass vase', 'quantity': 1, 'unit_price': 1}], status='draft')
        ids = lambda qs: [q['id'] for q in self.c.get('/api/quotes?' + qs)[1]['quotes']]
        self.assertEqual(ids('customer_id=%d' % a), [q1])
        self.assertEqual(ids('status=draft&search=Beta'), [q2])
        self.assertEqual(ids('search=oak'), [q1])                          # 按明细里的产品名/SKU 也能搜到
        self.assertEqual(ids('search=ZZ-OAK'), [q1])
        self.assertEqual(ids('search=Ann'), [q1])
        self.assertEqual(self.c.get('/api/quotes?search=%25')[1]['total'], 0)                     # % 当字面量
        self.assertEqual(self.c.get('/api/quotes?status=bogus')[0], 400)
        self.assertEqual(self.c.get('/api/quotes?customer_id=x')[0], 400)
        d = self.c.get('/api/quotes?limit=1&offset=0')[1]
        self.assertEqual(len(d['quotes']), 1)
        self.assertGreaterEqual(d['total'], 2)
        row = [q for q in self.c.get('/api/quotes?customer_id=%d' % a)[1]['quotes']][0]
        self.assertEqual((row['company'], row['item_count'], row['status_label'], row['past_validity']), ('Alpha Imports', 1, '已发送', False))

    def test_past_validity_flag_for_old_sent_quotes_only(self):
        qid, _ = self.mk_quote(valid_days=10)
        self.ctx.db.execute("UPDATE quotes SET created_at='2020-01-01 10:00:00' WHERE id=?", (qid,))
        q = self.quote(qid)
        self.assertEqual((q['valid_until'], q['past_validity'], q['status']), ('2020-01-11', True, 'sent'))         # 只是提示，状态不自动变
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'rejected'})
        self.assertFalse(self.quote(qid)['past_validity'])


class TestQuoteStatusEffects(QuoteBase):
    def test_stage_promotion_on_send_not_on_draft_and_never_regresses(self):
        for start, expect in (('', '已报价'), ('潜在', '已报价'), ('已联系', '已报价'), ('沉睡', '已报价'), ('已寄样', '已寄样'), ('成交', '成交'), ('复购', '复购')):
            cid = self.new_customer(stage=start)
            self.mk_quote(cid)
            self.assertEqual(self.stage(cid), expect, start)
        cid = self.new_customer(stage='潜在')
        qid, _ = self.mk_quote(cid, status='draft')
        self.assertEqual(self.stage(cid), '潜在')                                          # 草稿不算已报价
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'sent'})
        self.assertEqual(self.stage(cid), '已报价')

    def test_accept_writes_sell_history_once_with_customer_and_source(self):
        p1, p2 = self.new_product(sku='SH-1', cost=5, cost_currency='USD'), self.new_product(sku='SH-2', cost=5, cost_currency='USD')
        qid, cid = self.mk_quote(items=[{'product_id': p1, 'quantity': 10, 'unit_price': 8.8}, {'product_id': p2, 'quantity': 1, 'unit_price': 9.9},
                                        {'name': 'freight', 'quantity': 1, 'unit_price': 50}], currency='EUR')
        no = self.quote(qid)['quote_no']
        for _ in range(3):                                                                  # 反复点"成交"不能重复写
            self.c.post('/api/quotes/%d/status' % qid, {'status': 'accepted'})
        rows = self.sells()
        mine = [r for r in rows if r['source'] == no]
        self.assertEqual(sorted((r['product_id'], r['price'], r['currency'], r['customer_id']) for r in mine), sorted([(p1, 8.8, 'EUR', cid), (p2, 9.9, 'EUR', cid)]))
        self.assertEqual(mine[0]['effective_date'], time.strftime('%Y-%m-%d'))
        self.assertEqual(self.stage(cid), '成交')
        hints = self.c.get('/api/products/%d/last_price?customer_id=%d' % (p1, cid))[1]
        self.assertEqual((hints['to_this_customer']['price'], hints['to_this_customer']['source']), (8.8, no))     # 谈判提示用得上

    def test_second_accepted_quote_makes_repeat_customer_and_toggling_is_stable(self):
        cid = self.new_customer(stage='')
        q1, _ = self.mk_quote(cid)
        q2, _ = self.mk_quote(cid)
        self.c.post('/api/quotes/%d/status' % q1, {'status': 'accepted'})
        self.assertEqual(self.stage(cid), '成交')
        self.c.post('/api/quotes/%d/status' % q1, {'status': 'rejected'})
        self.c.post('/api/quotes/%d/status' % q1, {'status': 'accepted'})
        self.assertEqual(self.stage(cid), '成交')                                          # 同一张反复切换不会升级成复购
        self.c.post('/api/quotes/%d/status' % q2, {'status': 'accepted'})
        self.assertEqual(self.stage(cid), '复购')

    def test_unaccept_and_delete_revoke_history_and_realign(self):
        pid = self.new_product(sku='UN-1', cost=10, cost_currency='USD', profit_rate=0.5)
        cid = self.new_customer()
        qid, _ = self.mk_quote(cid, items=[{'product_id': pid, 'quantity': 1, 'unit_price': 99}])
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'accepted'})
        self.assertEqual(len(self.sells(pid)), 1)
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'sent'})
        self.assertEqual(self.sells(pid), [])                                               # 取消成交 = 撤销这批成交价
        self.assertEqual(self.c.get('/api/products/%d' % pid)[1]['product']['suggested_price'], 15.0)
        # 别的客户的成交价、手动录入的售价不受影响
        other = self.new_customer()
        self.c.post('/api/products/%d/price_history' % pid, {'price_type': 'sell', 'price': 12, 'currency': 'USD', 'customer_id': other, 'source': 'manual'})
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'accepted'})
        self.assertEqual(len(self.sells(pid)), 2)
        st, r = self.c.delete('/api/quotes/%d' % qid)
        self.assertEqual((st, r['needs_confirm']), (409, True))                             # 已成交的删除要二次确认
        self.assertEqual(self.c.delete('/api/quotes/%d' % qid, {'confirm': True})[0], 200)
        self.assertEqual([r['source'] for r in self.sells(pid)], ['manual'])
        self.assertEqual(self.ctx.db.scalar('SELECT COUNT(*) FROM quote_items WHERE quote_id=?', (qid,)), 0)
        self.assertEqual(self.c.get('/api/quotes/%d' % qid)[0], 404)
        self.assertEqual(self.c.delete('/api/quotes/%d' % qid, {'confirm': True})[0], 404)

    def test_delete_non_accepted_needs_no_confirm_and_status_validation(self):
        qid, cid = self.mk_quote()
        self.assertEqual(self.c.post('/api/quotes/%d/status' % qid, {'status': 'bogus'})[0], 400)
        self.assertEqual(self.c.post('/api/quotes/987654/status', {'status': 'sent'})[0], 404)
        self.assertEqual(self.c.delete('/api/quotes/%d' % qid)[0], 200)
        self.assertTrue(any('已删除' in n['content'] for n in self.c.get('/api/customers/%d' % cid)[1]['notes']))

    def test_feedback_saved_and_timeline_notes(self):
        qid, cid = self.mk_quote()
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'rejected', 'feedback': '价格太高，对手 $2.1'})
        q = self.quote(qid)
        self.assertEqual((q['status'], q['customer_feedback']), ('rejected', '价格太高，对手 $2.1'))
        self.assertTrue(any('未成交' in n['content'] and '价格太高' in n['content'] for n in self.c.get('/api/customers/%d' % cid)[1]['notes']))
        self.c.post('/api/quotes/%d/status' % qid, {'status': 'expired'})
        self.assertEqual(self.quote(qid)['customer_feedback'], '价格太高，对手 $2.1')       # 不传 feedback 不会清空

    def test_create_directly_as_accepted(self):
        pid = self.new_product(cost=1, cost_currency='USD')
        qid, cid = self.mk_quote(items=[{'product_id': pid, 'quantity': 1, 'unit_price': 3}], status='accepted')
        self.assertEqual((len(self.sells(pid)), self.stage(cid)), (1, '成交'))

    def test_deleting_customer_removes_quotes_but_keeps_price_history(self):
        pid = self.new_product(cost=1, cost_currency='USD')
        qid, cid = self.mk_quote(items=[{'product_id': pid, 'quantity': 1, 'unit_price': 3}], status='accepted')
        self.c.delete('/api/customers/%d' % cid, {'confirm': True})
        self.assertEqual(self.c.get('/api/quotes/%d' % qid)[0], 404)
        rows = self.sells(pid)
        self.assertEqual((len(rows), rows[0]['customer_id']), (1, None))


class TestWhatsApp(QuoteBase):
    def test_default_template_and_variables(self):
        self.c.put('/api/settings/quote', {'company_name': 'SINLUX Lighting'})
        qid, _ = self.mk_quote(self.new_customer(company='WA Buyer', name='Mia', whatsapp='+1 (555) 010-9999; +86 138'),
                               items=[{'sku': 'W-1', 'name': 'Lamp', 'spec': 'line1\n  line2   ' + 'x' * 100, 'quantity': 10, 'unit': 'pcs', 'unit_price': 2.5}],
                               lead_time='25 days', payment_terms='T/T 30%', valid_days=14)
        r = self.c.get('/api/quotes/%d/whatsapp' % qid)[1]
        t = r['text']
        no = self.quote(qid)['quote_no']
        self.assertTrue(t.startswith('Hi Mia,'))
        self.assertIn('quotation %s' % no, t)
        self.assertIn('- W-1 Lamp (line1 line2 ' + 'x' * 68 + '…): 10 pcs × USD 2.5 = USD 25.00', t)       # 规格折叠空白并截断
        for s in ('Total: USD 25.00', 'Lead time: 25 days', 'Payment: T/T 30%', 'Validity: 14 days', 'Best regards,\nSINLUX Lighting'):
            self.assertIn(s, t)
        self.assertEqual(r['wa_number'], '15550109999')                                  # 取第一个号码，去掉符号，可拼 wa.me 链接

    def test_missing_terms_show_tbd_and_unknown_wa_number(self):
        qid, _ = self.mk_quote(self.new_customer(company='NoTerms', whatsapp='abc'))
        r = self.c.get('/api/quotes/%d/whatsapp' % qid)[1]
        self.assertIn('Lead time: TBD', r['text'])
        self.assertEqual(r['wa_number'], '')

    def test_custom_template_is_safe_against_format_tricks(self):
        self.c.put('/api/settings/quote', {'whatsapp_template': 'Hello {customer_name} {0} {company.__class__} {unknown} {{kept}} {quote_no}}'})
        qid, _ = self.mk_quote(self.new_customer(company='Tmpl Co', name='Tom'))
        t = self.c.get('/api/quotes/%d/whatsapp' % qid)[1]['text']
        no = self.quote(qid)['quote_no']
        self.assertEqual(t, 'Hello Tom {0} {company.__class__} {unknown} {{kept}} %s}' % no)       # 只替换已知变量，别的原样、不报错
        self.c.put('/api/settings/quote', {'whatsapp_template': ''})
        self.assertIn('Please find our quotation', self.c.get('/api/quotes/%d/whatsapp' % qid)[1]['text'])   # 清空 = 恢复默认

    def test_render_template_unit(self):
        self.assertEqual(render_template('a {x} b {y} {x}', {'x': 1}), 'a 1 b {y} 1')
        self.assertEqual(render_template('{ }{}{{x}}', {'x': 'v'}), '{ }{}{v}')
        self.assertIn('{sender}', DEFAULT_WA_TEMPLATE)


class TestQuoteSettings(QuoteBase):
    def test_company_info_roundtrip_and_validation(self):
        d = self.c.get('/api/settings/quote')[1]
        self.assertEqual((d['company_name'], d['company_name_effective'], d['whatsapp_template_effective'] == DEFAULT_WA_TEMPLATE), ('', 'SINLUX', True))
        self.assertIn('customer_name', d['variables'])
        r = self.c.put('/api/settings/quote', {'company_name': ' Caribe FF&E ', 'company_email': 'sales@caribe.com', 'company_phone': '+1 555', 'company_address': 'Kingston'})[1]
        self.assertEqual((r['company_name'], r['company_email']), ('Caribe FF&E', 'sales@caribe.com'))
        for bad in ({'company_email': 'not-an-email'}, {'company_name': 'x' * 101}, {'whatsapp_template': 'x' * 4001}):
            self.assertEqual(self.c.put('/api/settings/quote', bad)[0], 400, bad)
        self.assertEqual(self.c.get('/api/settings/quote')[1]['company_email'], 'sales@caribe.com')
        self.assertEqual(self.c.put('/api/settings/quote', {'evil_key': 'x'})[0], 200)
        self.assertEqual(self.ctx.db.get_setting('evil_key'), '')


class TestQuoteFiles(QuoteBase):
    def setUp(self):
        self.c.put('/api/settings/quote', {'company_name': 'SINLUX Test Co', 'company_email': 'sales@sinlux.test', 'company_phone': '+86 755 0000', 'company_address': 'Shenzhen, China'})

    def export(self, qid, kind):
        st, r = self.c.post('/api/quotes/%d/%s' % (qid, kind))
        self.assertEqual(st, 200, r)
        st, raw, h = self.c.call('GET', r['url'])
        self.assertEqual(st, 200)
        self.assertIn('attachment', h['Content-Disposition'])
        return r, raw, h

    def test_pdf_content_chinese_images_and_layout(self):
        from imgutil import PNG_RED_BOX
        st, up = self.c.j('POST', '/api/images/upload', raw=PNG_RED_BOX)
        pid = self.new_product(sku='PDF-1', name='玻璃花瓶 Glass Vase', images=[up['image']['id']])
        cid = self.new_customer(company='Hotel Playa 酒店', name='Carlos', country='Jamaica', address='12 Beach Rd, Montego Bay')
        qid, _ = self.mk_quote(cid, items=[{'product_id': pid, 'sku': 'PDF-1', 'name': '玻璃花瓶 Glass Vase', 'spec': '高 30cm / 手工吹制 <b>bold?</b> & "quotes"', 'quantity': 120, 'unit': 'pcs', 'unit_price': 12.345, 'remark': 'Color: amber'},
                                           {'sku': 'X-2', 'name': 'Custom rug', 'quantity': 2.5, 'unit': 'm²', 'unit_price': 40}],
                               lead_time='30-35 days', payment_terms='T/T 30% deposit', shipping_terms='FOB Shenzhen', notes='Prices valid for 30 days.', valid_days=30)
        r, raw, h = self.export(qid, 'pdf')
        no = self.quote(qid)['quote_no']
        self.assertEqual((r['filename'], h['Content-Type']), (no + '.pdf', 'application/pdf'))
        self.assertTrue(raw.startswith(b'%PDF'))
        text = pdf_text(raw)
        for s in ('SINLUX Test Co', 'sales@sinlux.test', 'QUOTATION', no, 'Hotel Playa 酒店', 'Attn: Carlos', 'Jamaica', 'Montego Bay', 'PDF-1', '玻璃花瓶', 'Glass Vase',
                  'Color: amber', 'T/T 30% deposit', 'FOB Shenzhen', 'Prices valid for 30 days.', 'Thank you for your business!', 'Page 1 / 1'):
            self.assertIn(s, text, s)
        self.assertIn('<b>bold?</b> & "quotes"', text)                                      # 用户文本里的尖括号/&号原样出现，没被当成标记
        self.assertIn('1,481.40', text)                                                     # 120 × 12.345 = 1481.40
        self.assertIn('100.00', text)
        self.assertIn('1,581.40', text)                                                     # 合计
        self.assertEqual(pdf_pages(raw), 1)
        with tempfile.TemporaryDirectory() as d:                                            # 产品图真的画进去了（PDF 里有图片对象）
            p = os.path.join(d, 'q.pdf'); open(p, 'wb').write(raw)
            imgs = subprocess.run(['pdfimages', '-list', p], capture_output=True, text=True).stdout
            self.assertGreaterEqual(len(imgs.strip().splitlines()), 3)

    def test_pdf_many_rows_paginate_and_huge_text_does_not_crash(self):
        items = [{'sku': 'R-%03d' % i, 'name': 'Item number %d with a reasonably long descriptive name' % i, 'spec': 'spec line ' * 30, 'quantity': i + 1, 'unit_price': 1.5} for i in range(120)]
        items[3]['spec'] = 'very long ' * 390                                                # 超长规格：截断后仍能出页
        items[4]['name'] = '名' * 300
        qid, _ = self.mk_quote(items=items)
        r, raw, h = self.export(qid, 'pdf')
        pages = pdf_pages(raw)
        self.assertGreater(pages, 5)
        text = pdf_text(raw)
        self.assertIn('Page %d / %d' % (pages, pages), text)
        self.assertIn('R-119', text)
        self.assertIn('Page 1 /', text)
        self.assertEqual(len(re.findall(r'^\s*QUOTATION', text, re.M)), 1)                  # 抬头只在第一页

    def test_pdf_without_company_settings_and_without_images(self):
        self.c.put('/api/settings/quote', {'company_name': '', 'company_email': '', 'company_phone': '', 'company_address': ''})
        qid, _ = self.mk_quote()
        text = pdf_text(self.export(qid, 'pdf')[1])
        self.assertIn('SINLUX', text)                                                       # 默认抬头
        self.assertIn('Lamp', text)

    def test_excel_formulas_rounding_and_texts(self):
        qid, cid = self.mk_quote(self.new_customer(company='=cmd|evil Co', name='Eve'), items=[
            {'sku': 'E-1', 'name': 'Alpha', 'spec': '=HYPERLINK("http://x","y")', 'quantity': 3, 'unit': 'pcs', 'unit_price': 0.335},
            {'sku': 'E-2', 'name': '=1+1', 'quantity': 7, 'unit_price': 1.005}], notes='=SUM(A1)', payment_terms='T/T')
        r, raw, h = self.export(qid, 'excel')
        self.assertIn('spreadsheetml', h['Content-Type'])
        ws = load_workbook(io.BytesIO(raw)).active
        cells = {c.coordinate: c for row in ws.iter_rows() for c in row if c.value is not None}
        find = lambda text: [c for c in cells.values() if c.value == text]
        self.assertEqual(cells['A1'].value, 'SINLUX Test Co')
        hdr = [c.value for c in ws[9]]
        self.assertEqual(hdr[:9], ['#', 'Image', 'SKU', 'Description', 'Specification', 'Qty', 'Unit', 'Price (USD)', 'Amount (USD)'])
        self.assertEqual((ws['I10'].value, ws['I11'].value, ws['I12'].value), ('=ROUND(F10*H10,2)', '=ROUND(F11*H11,2)', '=SUM(I10:I11)'))
        self.assertEqual((ws['F10'].value, ws['H10'].value, ws['H11'].value), (3, 0.335, 1.005))
        self.assertEqual(self.quote(qid)['total'], 8.05)                                    # 1.01 + 7.04（ROUND 后）
        for text in ('=HYPERLINK("http://x","y")', '=1+1', '=cmd|evil Co', '=SUM(A1)'):    # 用户文本以 = 开头：存成文本，不会被 Excel 当公式执行
            hits = [c for c in cells.values() if isinstance(c.value, str) and text in c.value]
            self.assertTrue(hits, text)
            self.assertTrue(all(c.data_type == 's' for c in hits), text)

    def test_excel_embeds_product_thumbnails(self):
        from imgutil import PNG_RED_BOX
        up = self.c.j('POST', '/api/images/upload', raw=PNG_RED_BOX)[1]['image']
        pid = self.new_product(sku='XI-1', images=[up['id']])
        qid, _ = self.mk_quote(items=[{'product_id': pid, 'sku': 'XI-1', 'name': 'With image', 'quantity': 1, 'unit_price': 1}, {'name': 'No image', 'quantity': 1, 'unit_price': 1}])
        ws = load_workbook(io.BytesIO(self.export(qid, 'excel')[1])).active
        self.assertEqual(len(ws._images), 1)

    def test_export_download_names_are_safe(self):
        for p in ('/exports/..%2fcrm.db', '/exports/x.txt', '/exports/missing.pdf'):
            self.assertEqual(self.c.call('GET', p)[0], 404, p)
        self.assertEqual(self.c.post('/api/quotes/987654/pdf')[0], 404)
        self.assertEqual(self.c.post('/api/quotes/987654/excel')[0], 404)

    def test_export_filename_sanitised_for_odd_quote_numbers(self):
        cid = self.new_customer()
        qid = self.ctx.db.execute("INSERT INTO quotes(quote_no,customer_id,total) VALUES(?,?,0)", ('../weird no/1', cid)).lastrowid
        self.ctx.db.execute("INSERT INTO quote_items(quote_id,sku,name,quantity,unit_price) VALUES(?,?,?,1,1)", (qid, 'a', 'b'))
        r = self.c.post('/api/quotes/%d/pdf' % qid)
        self.assertEqual(r[0], 200)
        self.assertEqual(os.path.dirname(os.path.join(self.ctx.exports_dir, r[1]['filename'])), self.ctx.exports_dir)
        self.assertNotIn('/', r[1]['filename'])


class TestLegacyQuotes(QuoteBase):
    legacy = staticmethod(lambda d: build_legacy_db(d, 6))

    def test_old_quotes_visible_editable_and_exportable(self):
        rows = self.c.get('/api/quotes?limit=50')[1]
        self.assertEqual(rows['total'], 3)
        old = [q for q in rows['quotes'] if q['quote_no'] == 'SLQ-20260301-001'][0]
        q = self.quote(old['id'])
        self.assertEqual((q['status'], q['status_label'], q['total'], q['items'][0]['amount']), ('accepted', '成交', 220.0, 220.0))
        r, raw, _ = (lambda r: (r, *self.c.call('GET', r['url'])[1:]))(self.c.post('/api/quotes/%d/pdf' % old['id'])[1])
        self.assertTrue(raw.startswith(b'%PDF'))
        self.assertIn('SLQ-20260301-001', pdf_text(raw))
        self.assertEqual(self.c.post('/api/quotes/%d/excel' % old['id'])[0], 200)
        self.assertIn('Total: USD 220.00', self.c.get('/api/quotes/%d/whatsapp' % old['id'])[1]['text'])
        # 旧库里的成交单取消成交：旧库的售价历史（来源 PI-001，不是这张报价单）不能被误删
        before = self.ctx.db.scalar("SELECT COUNT(*) FROM price_history WHERE source='PI-001'")
        self.c.post('/api/quotes/%d/status' % old['id'], {'status': 'sent'})
        self.assertEqual(self.ctx.db.scalar("SELECT COUNT(*) FROM price_history WHERE source='PI-001'"), before)
        # 新建的报价编号不与旧编号冲突
        qid, _ = self.mk_quote(4)
        self.assertTrue(self.quote(qid)['quote_no'].startswith('SLQ-%s-' % TODAY))


if __name__ == '__main__':
    unittest.main()
